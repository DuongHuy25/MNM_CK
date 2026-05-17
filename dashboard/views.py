from django.shortcuts import render, redirect, get_object_or_404
from django.views.decorators.http import require_http_methods
from django.contrib.auth import get_user_model, authenticate, login as auth_login, logout as auth_logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q, Sum, Count, Avg, F
from django.http import HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.utils import timezone

from rest_framework import viewsets, permissions, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from dashboard.models import (
    Store, Department, Employee, Category, Brand, Product,
    Warehouse, WarehouseItem, WarehouseTransaction,
    StockBalance, StockMovement, Supplier, PurchaseOrder, PurchaseOrderItem,
    GoodsReceipt, GoodsReceiptItem, CustomerGroup, Customer, PaymentMethod, Order,
    OrderItem, OrderPayment, WarehouseBatch, WarehouseBatchItem,
    About, CustomerProfile, News, Review, ReviewImage, OrderReview,
    ContactSettings, DiscountCode, DiscountCodeUsage
)
from .serializers import (
    ProductSerializer, StoreSerializer, OrderSerializer, CategorySerializer,
    EmployeeSerializer, SupplierSerializer, CustomerSerializer, WarehouseSerializer
)
from .forms import (
    StoreForm, ProductForm, OrderForm, CategoryForm, SearchForm, EmployeeForm,
    SupplierForm, PurchaseOrderForm, CustomerForm, WarehouseBatchForm, BrandForm,
    UserForm, WarehouseForm, WarehouseItemForm, WarehouseTransactionForm,
    WarehouseBatchItemForm, ImportExcelForm, AboutForm, AboutImportForm,
    CustomerProfileForm, UserBasicInfoForm, NewsForm, ReviewReplyForm,
    ContactSettingsForm, DiscountCodeForm
)
from .utils import admin_required, search_items, paginate_queryset, get_stats, generate_batch_number


# ==================== AUTH ====================
def login_view(request):
    """Đăng nhập dashboard"""
    if request.user.is_authenticated:
        return redirect('dashboard:index')
    if request.method == 'POST':
        email = request.POST.get('email', '')
        password = request.POST.get('password', '')
        user = authenticate(request, username=email, password=password)
        if user is not None:
            auth_login(request, user)
            next_url = request.GET.get('next', '')
            return redirect(next_url if next_url else 'dashboard:index')
        messages.error(request, 'Email hoặc mật khẩu không đúng.')
    return render(request, 'dashboard/login.html', {'page_title': 'Đăng nhập'})


def logout_view(request):
    """Đăng xuất dashboard"""
    auth_logout(request)
    return redirect('dashboard:login')


# ==================== DASHBOARD ====================
@admin_required
def dashboard_index(request):
    """Dashboard chính"""
    try:
        from datetime import datetime, timedelta
        from django.db.models import Count, Sum
        from django.utils import timezone
        import json

        stats = {
            'stores_count': Store.objects.filter(is_active=True).count(),
            'products_count': Product.objects.filter(is_active=True).count(),
            'pending_orders': Order.objects.filter(status='pending').count(),
            'customers_count': Customer.objects.filter(is_active=True).count(),
            'employees_count': Employee.objects.filter(is_active=True).count(),
            'low_stock_count': StockBalance.objects.filter(quantity_on_hand__lte=10).count(),
        }

        recent_orders = Order.objects.select_related('customer', 'store').order_by('-created_at')[:5]
        recent_products = Product.objects.select_related('category').order_by('-created_at')[:5]
        recent_employees = Employee.objects.select_related('store', 'department').order_by('-created_at')[:5]
        # Normalize map coordinates so Leaflet can always render markers
        stores = list(
            Store.objects.filter(is_active=True)
            .exclude(latitude__isnull=True)
            .exclude(longitude__isnull=True)
            .values('id', 'name', 'address', 'latitude', 'longitude')
        )
        # Force lat/lng to float (in case DB stores as string)
        for s in stores:
            try:
                s['latitude'] = float(s['latitude'])
            except (TypeError, ValueError):
                s['latitude'] = None
            try:
                s['longitude'] = float(s['longitude'])
            except (TypeError, ValueError):
                s['longitude'] = None
        stores = [s for s in stores if s['latitude'] is not None and s['longitude'] is not None]


        # Chart data - Revenue by month (last 6 months)
        revenue_chart_data = []
        order_chart_data = []
        
        for i in range(6):
            month_date = timezone.now() - timedelta(days=30*i)
            month_start = month_date.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
            month_end = (month_start + timedelta(days=32)).replace(day=1, hour=23, minute=59, second=59, microsecond=999999) - timedelta(microseconds=1)
            
            month_revenue = Order.objects.filter(
                created_at__gte=month_start,
                created_at__lte=month_end,
                status='completed'
            ).aggregate(total=Sum('total_amount'))['total'] or 0
            
            month_orders = Order.objects.filter(
                created_at__gte=month_start,
                created_at__lte=month_end
            ).count()
            
            revenue_chart_data.append({
                'month': month_start.strftime('%m/%Y'),
                'revenue': float(month_revenue)
            })
            
            order_chart_data.append({
                'month': month_start.strftime('%m/%Y'),
                'orders': month_orders
            })
        
        revenue_chart_data.reverse()
        order_chart_data.reverse()

        # Top selling products
        top_products = OrderItem.objects.values(
            'product__name'
        ).annotate(
            total_quantity=Sum('quantity'),
            total_revenue=Sum('total_price')
        ).order_by('-total_quantity')[:10]

        # Orders by status
        orders_by_status = Order.objects.values('status').annotate(
            count=Count('id')
        ).order_by('status')

        # Low stock alerts
        low_stock_items = StockBalance.objects.filter(
            quantity_on_hand__lte=F('product__min_stock')
        ).select_related('product', 'warehouse')[:10]

        context = {
            'stats': stats,
            'recent_orders': recent_orders,
            'recent_products': recent_products,
            'recent_employees': recent_employees,
            'stores': stores,
            'revenue_chart_data': json.dumps(revenue_chart_data),
            'order_chart_data': json.dumps(order_chart_data),
            'top_products': top_products,
            'orders_by_status': orders_by_status,
            'low_stock_items': low_stock_items,
            'page_title': 'Bảng Điều Khiển',
        }
        return render(request, 'dashboard/index.html', context)
    except Exception as e:
        messages.error(request, f'Lỗi: {str(e)}')
        return render(request, 'dashboard/index.html', {
            'page_title': 'Bảng Điều Khiển',
            'stores': [],
        })


# ==================== STORES ====================
@admin_required
def stores_list(request):
    stores = Store.objects.all()
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        stores = stores.filter(Q(name__icontains=q) | Q(address__icontains=q))
    page_obj, paginator = paginate_queryset(stores, request.GET.get('page'), 15)
    return render(request, 'dashboard/stores/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Cửa Hàng',
    })


@admin_required
def stores_detail(request, pk):
    store = get_object_or_404(Store, pk=pk)
    employees = store.employees.all()
    orders = store.orders.all().order_by('-created_at')[:10]
    return render(request, 'dashboard/stores/detail.html', {
        'store': store, 'employees': employees, 'orders': orders,
        'page_title': f'Chi tiết: {store.name}',
    })


@admin_required
def stores_create(request):
    if request.method == 'POST':
        form = StoreForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cửa hàng đã được tạo thành công!')
            return redirect('dashboard:stores_list')
    else:
        form = StoreForm()
    return render(request, 'dashboard/stores/form.html', {
        'form': form, 'page_title': 'Thêm Cửa Hàng', 'is_create': True,
    })


@admin_required
def stores_edit(request, pk):
    store = get_object_or_404(Store, pk=pk)
    if request.method == 'POST':
        form = StoreForm(request.POST, instance=store)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cửa hàng đã được cập nhật!')
            return redirect('dashboard:stores_list')
    else:
        form = StoreForm(instance=store)
    return render(request, 'dashboard/stores/form.html', {
        'form': form, 'store': store, 'page_title': f'Chỉnh sửa: {store.name}', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def stores_delete(request, pk):
    store = get_object_or_404(Store, pk=pk)
    name = store.name
    store.delete()
    messages.success(request, f'Cửa hàng "{name}" đã được xóa!')
    return redirect('dashboard:stores_list')


# ==================== EMPLOYEES ====================
@admin_required
def employees_list(request):
    employees = Employee.objects.select_related('store', 'department')
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        employees = employees.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q) | 
            Q(email__icontains=q) | Q(phone__icontains=q)
        )
    page_obj, paginator = paginate_queryset(employees, request.GET.get('page'), 15)
    return render(request, 'dashboard/employees/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Nhân Viên',
    })


@admin_required
def employees_create(request):
    if request.method == 'POST':
        form = EmployeeForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Nhân viên đã được thêm thành công!')
            return redirect('dashboard:employees_list')
    else:
        form = EmployeeForm()
    return render(request, 'dashboard/employees/form.html', {
        'form': form, 'page_title': 'Thêm Nhân Viên', 'is_create': True,
    })


@admin_required
def employees_detail(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    return render(request, 'dashboard/employees/detail.html', {
        'employee': employee, 'page_title': f'Chi tiết: {employee.first_name} {employee.last_name}',
    })


@admin_required
def employees_edit(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == 'POST':
        form = EmployeeForm(request.POST, instance=employee)
        if form.is_valid():
            form.save()
            messages.success(request, 'Nhân viên đã được cập nhật!')
            return redirect('dashboard:employees_list')
    else:
        form = EmployeeForm(instance=employee)
    return render(request, 'dashboard/employees/form.html', {
        'form': form, 'employee': employee, 'page_title': f'Chỉnh sửa: {employee.first_name}', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def employees_delete(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    name = f"{employee.first_name} {employee.last_name}"
    employee.delete()
    messages.success(request, f'Nhân viên "{name}" đã được xóa!')
    return redirect('dashboard:employees_list')


# ==================== BRANDS ====================
@admin_required
def brands_list(request):
    brands = Brand.objects.annotate(product_count=Count('products'))
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        brands = brands.filter(Q(name__icontains=q) | Q(description__icontains=q))
    page_obj, paginator = paginate_queryset(brands, request.GET.get('page'), 15)
    return render(request, 'dashboard/brands/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Nhãn Hiệu',
    })


@admin_required
def brands_detail(request, pk):
    brand = get_object_or_404(Brand, pk=pk)
    products = Product.objects.filter(brand=brand).select_related('category')
    page_obj, paginator = paginate_queryset(products.order_by('-created_at'), request.GET.get('page'), 20)
    return render(request, 'dashboard/brands/detail.html', {
        'brand': brand, 'page_obj': page_obj, 'paginator': paginator,
        'page_title': f'Nhãn Hiệu - {brand.name}',
    })


@admin_required
def brands_create(request):
    if request.method == 'POST':
        form = BrandForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Nhãn hiệu đã được tạo!')
            return redirect('dashboard:brands_list')
    else:
        form = BrandForm()
    return render(request, 'dashboard/brands/form.html', {
        'form': form, 'page_title': 'Thêm Nhãn Hiệu', 'is_create': True,
    })


@admin_required
def brands_edit(request, pk):
    brand = get_object_or_404(Brand, pk=pk)
    if request.method == 'POST':
        form = BrandForm(request.POST, instance=brand)
        if form.is_valid():
            form.save()
            messages.success(request, 'Nhãn hiệu đã được cập nhật!')
            return redirect('dashboard:brands_list')
    else:
        form = BrandForm(instance=brand)
    return render(request, 'dashboard/brands/form.html', {
        'form': form, 'brand': brand,
        'page_title': f'Chỉnh Sửa: {brand.name}', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def brands_delete(request, pk):
    brand = get_object_or_404(Brand, pk=pk)
    name = brand.name
    brand.delete()
    messages.success(request, f'Nhãn hiệu "{name}" đã được xóa!')
    return redirect('dashboard:brands_list')


# ==================== CATEGORIES ====================
@admin_required
def categories_list(request):
    categories = Category.objects.all()
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        categories = categories.filter(Q(name__icontains=q) | Q(slug__icontains=q))
    page_obj, paginator = paginate_queryset(categories, request.GET.get('page'), 15)
    return render(request, 'dashboard/categories/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Danh Mục',
    })


@admin_required
def categories_create(request):
    if request.method == 'POST':
        form = CategoryForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Danh mục đã được tạo!')
            return redirect('dashboard:categories_list')
    else:
        form = CategoryForm()
    return render(request, 'dashboard/categories/form.html', {
        'form': form, 'page_title': 'Thêm Danh Mục', 'is_create': True,
    })


@admin_required
def categories_edit(request, pk):
    category = get_object_or_404(Category, pk=pk)
    if request.method == 'POST':
        form = CategoryForm(request.POST, request.FILES, instance=category)
        if form.is_valid():
            form.save()
            messages.success(request, 'Danh mục đã được cập nhật!')
            return redirect('dashboard:categories_list')
    else:
        form = CategoryForm(instance=category)
    return render(request, 'dashboard/categories/form.html', {
        'form': form, 'category': category,
        'page_title': f'Chỉnh sửa: {category.name}', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def categories_delete(request, pk):
    category = get_object_or_404(Category, pk=pk)
    name = category.name
    category.delete()
    messages.success(request, f'Danh mục "{name}" đã được xóa!')
    return redirect('dashboard:categories_list')


# ==================== PRODUCTS ====================
@admin_required
def products_list(request):
    products = Product.objects.select_related('category', 'brand')
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        products = products.filter(
            Q(name__icontains=q) | Q(description__icontains=q) | Q(barcode__icontains=q)
        )
    category_filter = request.GET.get('category')
    if category_filter:
        products = products.filter(category__id=category_filter)
    page_obj, paginator = paginate_queryset(products, request.GET.get('page'), 15)
    categories = Category.objects.all()
    return render(request, 'dashboard/products/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'categories': categories,
        'current_category': category_filter, 'page_title': 'Quản Lý Sản Phẩm',
    })


@admin_required
def products_detail(request, pk):
    product = get_object_or_404(Product.objects.select_related('category', 'brand'), pk=pk)
    stock_balances = WarehouseItem.objects.filter(product=product).select_related('warehouse__store')
    return render(request, 'dashboard/products/detail.html', {
        'product': product, 'stock_balances': stock_balances,
        'page_title': f'Chi tiết: {product.name}',
    })


@admin_required
def products_create(request):
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, 'Sản phẩm đã được thêm thành công!')
            return redirect('dashboard:products_list')
    else:
        form = ProductForm()
    return render(request, 'dashboard/products/form.html', {
        'form': form, 'page_title': 'Thêm Sản Phẩm', 'is_create': True,
    })


@admin_required
def products_edit(request, pk):
    product = get_object_or_404(Product, pk=pk)
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES, instance=product)
        if form.is_valid():
            form.save()
            messages.success(request, 'Sản phẩm đã được cập nhật!')
            return redirect('dashboard:products_list')
    else:
        form = ProductForm(instance=product)
    return render(request, 'dashboard/products/form.html', {
        'form': form, 'product': product,
        'page_title': f'Chỉnh sửa: {product.name}', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def products_delete(request, pk):
    product = get_object_or_404(Product, pk=pk)
    name = product.name
    product.delete()
    messages.success(request, f'Sản phẩm "{name}" đã được xóa!')
    return redirect('dashboard:products_list')


# ==================== CUSTOMERS ====================
@admin_required
def customers_list(request):
    customers = Customer.objects.select_related('group')
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        customers = customers.filter(
            Q(first_name__icontains=q) | Q(last_name__icontains=q) | 
            Q(phone__icontains=q) | Q(email__icontains=q)
        )
    page_obj, paginator = paginate_queryset(customers, request.GET.get('page'), 15)
    return render(request, 'dashboard/customers/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Khách Hàng',
    })


@admin_required
def customers_create(request):
    if request.method == 'POST':
        form = CustomerForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Khách hàng đã được thêm thành công!')
            return redirect('dashboard:customers_list')
    else:
        form = CustomerForm()
    return render(request, 'dashboard/customers/form.html', {
        'form': form, 'page_title': 'Thêm Khách Hàng', 'is_create': True,
    })


@admin_required
def customers_edit(request, pk):
    customer = get_object_or_404(Customer, pk=pk)
    if request.method == 'POST':
        form = CustomerForm(request.POST, instance=customer)
        if form.is_valid():
            form.save()
            messages.success(request, 'Khách hàng đã được cập nhật!')
            return redirect('dashboard:customers_list')
    else:
        form = CustomerForm(instance=customer)
    return render(request, 'dashboard/customers/form.html', {
        'form': form, 'customer': customer, 'page_title': f'Chỉnh sửa', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def customers_delete(request, pk):
    customer = get_object_or_404(Customer, pk=pk)
    name = f"{customer.first_name} {customer.last_name}"
    customer.delete()
    messages.success(request, f'Khách hàng "{name}" đã được xóa!')
    return redirect('dashboard:customers_list')


# ==================== ORDERS ====================
@admin_required
def orders_list(request):
    orders = Order.objects.select_related('customer', 'store')
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        orders = orders.filter(Q(order_number__icontains=q) | Q(customer__phone__icontains=q))
    status_filter = request.GET.get('status')
    if status_filter:
        orders = orders.filter(status=status_filter)
    page_obj, paginator = paginate_queryset(orders.order_by('-created_at'), request.GET.get('page'), 15)
    return render(request, 'dashboard/orders/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Đơn Hàng',
    })


@admin_required
def orders_create(request):
    if request.method == 'POST':
        form = OrderForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Đơn hàng đã được tạo!')
            return redirect('dashboard:orders_list')
    else:
        form = OrderForm()
    return render(request, 'dashboard/orders/form.html', {
        'form': form, 'page_title': 'Thêm Đơn Hàng', 'is_create': True,
    })


@admin_required
def orders_detail(request, pk):
    order = get_object_or_404(Order, pk=pk)
    items = order.items.select_related('product')
    return render(request, 'dashboard/orders/detail.html', {
        'order': order, 'items': items, 'page_title': f'Chi tiết: {order.order_number}',
    })


@admin_required
def orders_edit(request, pk):
    order = get_object_or_404(Order, pk=pk)
    if request.method == 'POST':
        form = OrderForm(request.POST, instance=order)
        if form.is_valid():
            form.save()
            messages.success(request, 'Đơn hàng đã được cập nhật!')
            return redirect('dashboard:orders_list')
    else:
        form = OrderForm(instance=order)
    return render(request, 'dashboard/orders/form.html', {
        'form': form, 'order': order, 'page_title': f'Chỉnh sửa', 'is_edit': True,
    })


@admin_required
@require_http_methods(["POST"])
def orders_delete(request, pk):
    order = get_object_or_404(Order, pk=pk)
    order.delete()
    messages.success(request, f'Đơn hàng đã được xóa!')
    return redirect('dashboard:orders_list')


# ==================== WAREHOUSE ====================
@admin_required
def warehouse_list(request):
    # Stats
    total_warehouses = Warehouse.objects.count()
    total_items = WarehouseItem.objects.count()
    low_stock_items = WarehouseItem.objects.filter(quantity__lte=F('min_quantity')).select_related('product', 'warehouse__store')
    total_value = sum(wi.quantity * wi.unit_price for wi in WarehouseItem.objects.all())

    # Warehouse items with pagination
    warehouse_items = WarehouseItem.objects.select_related('product__brand', 'product__category', 'warehouse__store').order_by('-updated_at')
    page_obj, paginator = paginate_queryset(warehouse_items, request.GET.get('page'), 15)

    # Recent batches
    recent_batches = WarehouseBatch.objects.select_related('warehouse__store').order_by('-created_at')[:10]

    stats = {
        'total_warehouses': total_warehouses,
        'total_items': total_items,
        'low_stock_count': low_stock_items.count(),
        'total_value': total_value,
    }

    return render(request, 'dashboard/warehouse/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'low_stock_items': low_stock_items[:10],
        'recent_batches': recent_batches,
        'stats': stats, 'page_title': 'Quản Lý Kho',
    })


@admin_required
def warehouse_detail(request, pk):
    warehouse = get_object_or_404(Warehouse, pk=pk)
    stock_balances = warehouse.stock_balances.select_related('product')
    return render(request, 'dashboard/warehouse/detail.html', {
        'warehouse': warehouse, 'stock_balances': stock_balances,
        'page_title': f'Kho - {warehouse.store.name}',
    })


# ==================== REST FRAMEWORK ====================

class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.all()
    serializer_class = ProductSerializer
    permission_classes = [permissions.IsAuthenticated]


class StoreViewSet(viewsets.ModelViewSet):
    queryset = Store.objects.all()
    serializer_class = StoreSerializer
    permission_classes = [permissions.IsAuthenticated]


class CategoryViewSet(viewsets.ModelViewSet):
    queryset = Category.objects.all()
    serializer_class = CategorySerializer
    permission_classes = [permissions.IsAuthenticated]


class OrderViewSet(viewsets.ModelViewSet):
    queryset = Order.objects.all()
    serializer_class = OrderSerializer
    permission_classes = [permissions.IsAuthenticated]


@admin_required
def manage_stores_view(request):
    """View để quản lý cửa hàng toàn bộ"""
    stores = Store.objects.all()
    context = {
        'stores': stores,
        'page_title': 'Quản Lý Toàn Bộ Cửa Hàng',
    }
    return render(request, 'dashboard/manage_stores.html', context)


# API Views
@api_view(['GET'])
def api_search_stores(request):
    """API tìm kiếm cửa hàng"""
    q = request.GET.get('q', '')
    stores = Store.objects.filter(
        Q(name__icontains=q) | Q(address__icontains=q)
    )[:10]
    serializer = StoreSerializer(stores, many=True)
    return Response(serializer.data)


@api_view(['GET'])
def api_search_by_alley(request):
    """API tìm kiếm theo địa chỉ"""
    q = request.GET.get('q', '')
    stores = Store.objects.filter(address__icontains=q)[:10]
    serializer = StoreSerializer(stores, many=True)
    return Response(serializer.data)


@api_view(['GET'])
def api_warehouse_low_stock(request):
    """API lấy sản phẩm tồn kho thấp"""
    low_stock = StockBalance.objects.filter(
        quantity_on_hand__lte=10
    ).values_list('product__name', 'quantity_on_hand')
    return Response({'low_stock_items': list(low_stock)})


@api_view(['POST'])
def store_list_create(request):
    """API tạo cửa hàng"""
    if request.method == 'POST':
        serializer = StoreSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=201)
        return Response(serializer.errors, status=400)
    stores = Store.objects.all()
    serializer = StoreSerializer(stores, many=True)
    return Response(serializer.data)


# ==================== CHECKOUT API (merged from orders app) ====================

from django.db import transaction
from django.core.exceptions import ValidationError
from rest_framework.views import APIView
from .serializers import CheckoutSerializer


class CheckoutAPIView(APIView):
    """
    API checkout đơn hàng online.
    Sử dụng transaction.atomic() và select_for_update() để đảm bảo tính toàn vẹn dữ liệu.
    """
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request, *args, **kwargs):
        serializer = CheckoutSerializer(data=request.data)

        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

        validated_data = serializer.validated_data

        try:
            with transaction.atomic():
                total_amount = 0
                order_items = []

                # Validate all items and calculate total
                for item_data in validated_data['items']:
                    product_id = item_data['product_id']
                    quantity = item_data['quantity']

                    product = Product.objects.select_for_update().get(id=product_id)

                    if not product.is_active:
                        raise ValidationError(f"Product {product.name} is not available")

                    if product.stock < quantity:
                        raise ValidationError(
                            f"Not enough stock for {product.name}. "
                            f"Available: {product.stock}, Requested: {quantity}"
                        )

                    item_total = product.sale_price * quantity
                    total_amount += item_total

                # Create order with calculated total
                order = Order.objects.create(
                    user=request.user,
                    shipping_address=validated_data['shipping_address'],
                    phone_number=validated_data['phone_number'],
                    note=validated_data.get('note', ''),
                    total_amount=total_amount,
                    status='pending'
                )

                # Create order items and deduct stock
                for item_data in validated_data['items']:
                    product_id = item_data['product_id']
                    quantity = item_data['quantity']

                    product = Product.objects.select_for_update().get(id=product_id)

                    order_item = OrderItem.objects.create(
                        order=order,
                        product=product,
                        quantity=quantity,
                        unit_price=product.sale_price,
                        total_price=product.sale_price * quantity
                    )
                    order_items.append(order_item)

                    product.stock -= quantity
                    product.save()

                response_data = {
                    "message": "Order placed successfully",
                    "order_id": order.order_id,
                    "total_amount": float(total_amount),
                    "user_id": request.user.id,
                    "user_email": request.user.email,
                    "shipping_address": order.shipping_address,
                    "phone_number": order.phone_number,
                    "note": order.note,
                    "status": order.status,
                    "created_at": order.created_at.isoformat(),
                    "items": [
                        {
                            "product_id": item.product.id,
                            "product_name": item.product.name,
                            "quantity": item.quantity,
                            "price": float(item.unit_price),
                            "total_price": float(item.total_price)
                        }
                        for item in order_items
                    ]
                }

                return Response(response_data, status=status.HTTP_201_CREATED)

        except Product.DoesNotExist:
            return Response(
                {"error": "One or more products not found"},
                status=status.HTTP_400_BAD_REQUEST
            )
        except ValidationError as e:
            return Response(
                {"error": str(e)},
                status=status.HTTP_400_BAD_REQUEST
            )
        except Exception as e:
            return Response(
                {"error": "An error occurred while processing your order. Please try again."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def order_list_view(request):
    """API lấy danh sách đơn hàng của user."""
    orders = Order.objects.filter(user=request.user).order_by('-created_at')
    data = []
    for order in orders:
        data.append({
            "order_id": order.order_id,
            "total_amount": order.total_amount,
            "status": order.status,
            "created_at": order.created_at.isoformat(),
            "items_count": order.items.count()
        })
    return Response({"orders": data, "total_count": len(data)})


@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def order_detail_view(request, order_id):
    """API lấy chi tiết đơn hàng."""
    try:
        order = get_object_or_404(Order, order_id=order_id, user=request.user)
        data = {
            "order_id": order.order_id,
            "status": order.status,
            "total_amount": float(order.total_amount),
            "created_at": order.created_at.isoformat(),
            "items": [
                {
                    "product_name": item.product.name,
                    "quantity": item.quantity,
                    "price": float(item.price),
                    "total_price": float(item.total_price)
                }
                for item in order.items.all()
            ]
        }
        return Response(data)
    except Exception as e:
        return Response({"error": str(e)}, status=400)


# API endpoints for new CRUD views
@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def supplier_list_api(request):
    """API danh sách nhà cung cấp"""
    suppliers = Supplier.objects.all().values('id', 'name', 'contact_person', 'email', 'phone', 'is_active')
    return Response(list(suppliers))

@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def purchase_order_list_api(request):
    """API danh sách đơn mua hàng"""
    orders = PurchaseOrder.objects.select_related('supplier').all()
    data = []
    for order in orders:
        data.append({
            'id': order.pk,
            'po_number': order.po_number,
            'supplier': order.supplier.name,
            'order_date': order.order_date.isoformat(),
            'total_amount': float(order.total_amount),
            'status': order.status
        })
    return Response(data)

@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def goods_receipt_list_api(request):
    """API danh sách phiếu nhập hàng"""
    receipts = GoodsReceipt.objects.select_related('purchase_order', 'supplier').all()
    data = []
    for receipt in receipts:
        data.append({
            'id': receipt.pk,
            'receipt_number': receipt.receipt_number,
            'purchase_order': receipt.purchase_order.po_number if receipt.purchase_order else None,
            'supplier': receipt.supplier.name if receipt.supplier else None,
            'receipt_date': receipt.receipt_date.isoformat(),
            'total_amount': float(receipt.total_amount),
            'status': receipt.status
        })
    return Response(data)

@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def customer_group_list_api(request):
    """API danh sách nhóm khách hàng"""
    groups = CustomerGroup.objects.all().values('id', 'name', 'description', 'discount_percent', 'is_active')
    return Response(list(groups))

@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def stock_balance_list_api(request):
    """API danh sách số dư tồn kho"""
    balances = StockBalance.objects.select_related('product', 'warehouse').all()
    data = []
    for balance in balances:
        data.append({
            'id': balance.pk,
            'product': balance.product.name,
            'warehouse': balance.warehouse.store.name,
            'quantity_on_hand': balance.quantity_on_hand,
            'min_stock': balance.product.min_stock,
            'status': 'low_stock' if balance.quantity_on_hand <= balance.product.min_stock else 'normal'
        })
    return Response(data)

@api_view(['GET'])
@permission_classes([permissions.IsAuthenticated])
def stock_movement_list_api(request):
    """API danh sách chuyển động kho"""
    movements = StockMovement.objects.select_related('product', 'warehouse', 'created_by').all()[:50]
    data = []
    for movement in movements:
        creator = movement.created_by
        creator_label = ''
        if creator:
            user = getattr(creator, 'user', None)
            if user:
                creator_label = user.get_full_name() or user.email
            else:
                creator_label = f'{creator.first_name} {creator.last_name}'.strip() or str(creator.pk)
        data.append({
            'id': movement.pk,
            'product': movement.product.name,
            'warehouse': movement.warehouse.store.name,
            'movement_type': movement.movement_type,
            'quantity': movement.quantity,
            'note': movement.note,
            'created_at': movement.created_at.isoformat(),
            'created_by': creator_label,
        })
    return Response(data)


# ==================== PERMISSION CLASS (from dashboard1) ====================

class IsAdminOrReadOnly(permissions.BasePermission):
    def has_permission(self, request, view):
        if request.method in permissions.SAFE_METHODS:
            return request.user.is_authenticated
        return request.user and request.user.is_staff


# ==================== USERS (from dashboard1) ====================

User = get_user_model()

@admin_required
def users_list(request):
    users = User.objects.select_related('profile').all()
    search_form = SearchForm(request.GET or None)
    if search_form.is_valid() and search_form.cleaned_data.get('q'):
        q = search_form.cleaned_data['q']
        users = search_items(users, q, ['username', 'email', 'first_name', 'last_name'])
    page_obj, paginator = paginate_queryset(users, request.GET.get('page'), 6)
    return render(request, 'dashboard/users/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'search_form': search_form, 'page_title': 'Quản Lý Người Dùng',
    })

@admin_required
def users_create(request):
    if request.method == 'POST':
        form = UserForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            password = form.cleaned_data.get('password')
            if password:
                user.set_password(password)
            user.save()
            messages.success(request, f'Người dùng "{user.email}" đã được tạo thành công!')
            return redirect('dashboard:users_detail', pk=user.pk)
    else:
        form = UserForm()
    return render(request, 'dashboard/users/form.html', {
        'form': form, 'is_create': True,
        'page_title': 'Thêm Người Dùng',
    })

@admin_required
def users_detail(request, pk):
    user = get_object_or_404(User.objects.select_related('profile'), pk=pk)
    orders = Order.objects.filter(user=user).order_by('-created_at')[:10]
    return render(request, 'dashboard/users/detail.html', {
        'user_obj': user, 'orders': orders,
        'page_title': f'Chi tiết: {user.get_full_name() or user.email}',
    })

@admin_required
def users_edit(request, pk):
    user = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        form = UserForm(request.POST, instance=user)
        if form.is_valid():
            u = form.save(commit=False)
            password = form.cleaned_data.get('password')
            if password:
                u.set_password(password)
            u.save()
            messages.success(request, 'Người dùng đã được cập nhật!')
            return redirect('dashboard:users_detail', pk=user.pk)
    else:
        form = UserForm(instance=user)
    return render(request, 'dashboard/users/form.html', {
        'form': form, 'user_obj': user, 'is_create': False,
        'page_title': f'Chỉnh sửa: {user.get_full_name() or user.email}',
    })

@admin_required
@require_http_methods(["POST"])
def users_delete(request, pk):
    user = get_object_or_404(User, pk=pk)
    username = user.get_full_name() or user.email
    user.delete()
    messages.success(request, f'Người dùng "{username}" đã được xóa!')
    return redirect('dashboard:users_list')


# ==================== WAREHOUSE CRUD (from dashboard1) ====================

@admin_required
def warehouse_create(request):
    if request.method == 'POST':
        form = WarehouseForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Quản lý kho đã được thêm!')
            return redirect('dashboard:warehouse_list')
    else:
        form = WarehouseForm()
    return render(request, 'dashboard/warehouse/form.html', {
        'form': form, 'page_title': 'Thêm Quản Lý Kho', 'is_create': True,
    })

@admin_required
def warehouse_edit(request, pk):
    warehouse = get_object_or_404(Warehouse, pk=pk)
    if request.method == 'POST':
        form = WarehouseForm(request.POST, instance=warehouse)
        if form.is_valid():
            form.save()
            messages.success(request, 'Quản lý kho đã được cập nhật!')
            return redirect('dashboard:warehouse_detail', pk=warehouse.pk)
    else:
        form = WarehouseForm(instance=warehouse)
    return render(request, 'dashboard/warehouse/form.html', {
        'form': form, 'warehouse': warehouse,
        'page_title': f'Chỉnh Sửa Kho - {warehouse.store.name}', 'is_edit': True,
    })

@admin_required
@require_http_methods(["POST"])
def warehouse_delete(request, pk):
    warehouse = get_object_or_404(Warehouse, pk=pk)
    warehouse.delete()
    messages.success(request, 'Quản lý kho đã được xóa!')
    return redirect('dashboard:warehouse_list')


# ==================== WAREHOUSE ITEM (from dashboard1) ====================

@admin_required
def warehouse_item_create(request):
    if request.method == 'POST':
        form = WarehouseItemForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Sản phẩm đã được thêm vào kho!')
            return redirect('dashboard:warehouse_list')
    else:
        form = WarehouseItemForm()
    return render(request, 'dashboard/warehouse/form.html', {
        'form': form, 'page_title': 'Thêm Sản Phẩm Vào Kho', 'is_item': True, 'is_create': True,
    })

@admin_required
def warehouse_item_edit(request, pk):
    item = get_object_or_404(WarehouseItem.objects.select_related('product__brand', 'product__category'), pk=pk)
    product = item.product
    if request.method == 'POST':
        form = WarehouseItemForm(request.POST, instance=item)
        if form.is_valid():
            # Update warehouse item
            warehouse_item = form.save()
            # Also update linked product info
            product.name = request.POST.get('product_name', product.name)
            product.sale_price = float(request.POST.get('product_sale_price', product.sale_price) or product.sale_price)
            product.cost_price = float(request.POST.get('product_cost_price', product.cost_price) or product.cost_price)
            product.is_active = 'product_is_active' in request.POST
            product.save(update_fields=['name', 'sale_price', 'cost_price', 'is_active', 'updated_at'])
            messages.success(request, 'Sản phẩm và thông tin kho đã được cập nhật!')
            return redirect('dashboard:warehouse_list')
    else:
        form = WarehouseItemForm(instance=item)
    return render(request, 'dashboard/warehouse/form.html', {
        'form': form, 'warehouse_item': item, 'product': product,
        'page_title': 'Chỉnh Sửa Sản Phẩm', 'is_item': True, 'is_edit': True,
    })

@admin_required
@require_http_methods(["POST"])
def warehouse_item_delete(request, pk):
    item = get_object_or_404(WarehouseItem, pk=pk)
    item.delete()
    messages.success(request, 'Sản phẩm đã được xóa khỏi kho!')
    return redirect('dashboard:warehouse_list')


# ==================== WAREHOUSE TRANSACTION (from dashboard1) ====================

@admin_required
def warehouse_transaction_create(request):
    return redirect('dashboard:warehouse_batch_create')

@admin_required
@require_http_methods(["POST"])
def warehouse_transaction_delete(request, pk):
    transaction = get_object_or_404(WarehouseTransaction, pk=pk)
    item = transaction.warehouse_item
    if transaction.transaction_type == 'import':
        item.quantity -= transaction.quantity
    else:
        item.quantity += transaction.quantity
    item.save()
    transaction.delete()
    messages.success(request, 'Giao dịch đã được xóa!')
    return redirect('dashboard:warehouse_list')


# ==================== WAREHOUSE BATCH (from dashboard1) ====================

@admin_required
def warehouse_batch_list(request):
    batches = WarehouseBatch.objects.select_related('warehouse__store', 'created_by').order_by('-created_at')
    batch_type_filter = request.GET.get('batch_type')
    if batch_type_filter:
        batches = batches.filter(batch_type=batch_type_filter)
    page_obj, paginator = paginate_queryset(batches, request.GET.get('page'), 10)
    return render(request, 'dashboard/warehouse/batch_list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'page_title': 'Phiếu Nhập/Xuất Kho',
    })

@admin_required
def warehouse_batch_create(request):
    import openpyxl
    import csv
    import io

    if request.method == 'POST':
        action = request.POST.get('action', 'create')
        form = WarehouseBatchForm(request.POST)

        if action == 'import-excel':
            if 'excel_file' not in request.FILES:
                messages.error(request, 'Vui lòng chọn file Excel!')
                return render(request, 'dashboard/warehouse/batch_form.html', {
                    'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
                })
            if not form.is_valid():
                messages.error(request, f'Lỗi form: {form.errors}')
                return render(request, 'dashboard/warehouse/batch_form.html', {
                    'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
                })

            batch = form.save(commit=False)
            batch.created_by = request.user
            if not batch.batch_number:
                batch.batch_number = generate_batch_number(batch.batch_type)
            batch.save()

            try:
                excel_file = request.FILES['excel_file']
                filename = excel_file.name.lower()
                batch.items.all().delete()
                rows_added = 0
                errors = []

                if filename.endswith('.csv'):
                    file_content = io.TextIOWrapper(excel_file.file, encoding='utf-8')
                    reader = csv.reader(file_content)
                    next(reader)
                    for row_num, row in enumerate(reader, start=2):
                        if not row or not row[0].strip():
                            continue
                        try:
                            product_name = row[0].strip()
                            quantity = int(row[1]) if len(row) > 1 else 1
                            unit_price = float(row[2]) if len(row) > 2 else 0
                            product = Product.objects.filter(name__icontains=product_name).first()
                            if not product:
                                errors.append(f"Hàng {row_num}: Sản phẩm '{product_name}' không tìm thấy")
                                continue
                            WarehouseBatchItem.objects.create(
                                batch=batch, product=product, quantity=quantity, unit_price=unit_price
                            )
                            warehouse_item = WarehouseItem.objects.filter(
                                warehouse=batch.warehouse, product=product
                            ).first()
                            if warehouse_item:
                                if batch.batch_type == 'import':
                                    warehouse_item.quantity += quantity
                                else:
                                    warehouse_item.quantity -= quantity
                                warehouse_item.save()
                            rows_added += 1
                        except (ValueError, IndexError) as e:
                            errors.append(f"Hàng {row_num}: Lỗi dữ liệu - {str(e)}")

                elif filename.endswith(('.xlsx', '.xls')):
                    wb = openpyxl.load_workbook(excel_file)
                    ws = wb.active
                    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                        if not row or not row[0]:
                            continue
                        try:
                            product_name = str(row[0]).strip()
                            quantity = int(row[1]) if row[1] else 1
                            unit_price = float(row[2]) if row[2] else 0
                            product = Product.objects.filter(name__icontains=product_name).first()
                            if not product:
                                errors.append(f"Hàng {row_num}: Sản phẩm '{product_name}' không tìm thấy")
                                continue
                            WarehouseBatchItem.objects.create(
                                batch=batch, product=product, quantity=quantity, unit_price=unit_price
                            )
                            warehouse_item = WarehouseItem.objects.filter(
                                warehouse=batch.warehouse, product=product
                            ).first()
                            if warehouse_item:
                                if batch.batch_type == 'import':
                                    warehouse_item.quantity += quantity
                                else:
                                    warehouse_item.quantity -= quantity
                                warehouse_item.save()
                            rows_added += 1
                        except (ValueError, TypeError) as e:
                            errors.append(f"Hàng {row_num}: {str(e)}")
                else:
                    messages.error(request, 'Định dạng file không hỗ trợ. Vui lòng upload .xlsx, .xls hoặc .csv')
                    batch.delete()
                    return render(request, 'dashboard/warehouse/batch_form.html', {
                        'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
                    })

                batch.created_at = timezone.now()
                batch.calculate_total()
                batch.save()
                msg = f'Import thành công {rows_added} sản phẩm'
                if errors:
                    msg += f' ({len(errors)} lỗi)'
                    for error in errors[:5]:
                        messages.warning(request, error)
                messages.success(request, msg)
                return render(request, 'dashboard/warehouse/batch_form.html', {
                    'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
                })
            except Exception as e:
                messages.error(request, f'Lỗi: {str(e)}')
                batch.delete()
                return render(request, 'dashboard/warehouse/batch_form.html', {
                    'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
                })
        else:
            if form.is_valid():
                batch = form.save(commit=False)
                batch.created_by = request.user
                if not batch.batch_number:
                    batch.batch_number = generate_batch_number(batch.batch_type)
                batch.save()
                messages.success(request, 'Phiếu đã được tạo! Hãy thêm sản phẩm.')
                return redirect('dashboard:warehouse_batch_add_items', pk=batch.pk)
    else:
        form = WarehouseBatchForm()
    return render(request, 'dashboard/warehouse/batch_form.html', {
        'form': form, 'page_title': 'Tạo Phiếu Nhập/Xuất Kho', 'is_create': True,
    })

@admin_required
def warehouse_batch_add_items(request, pk):
    batch = get_object_or_404(WarehouseBatch, pk=pk)
    if request.method == 'POST':
        form = WarehouseBatchItemForm(request.POST)
        if form.is_valid():
            item = form.save(commit=False)
            item.batch = batch
            item.save()
            warehouse_item = WarehouseItem.objects.filter(
                warehouse=batch.warehouse, product=item.product
            ).first()
            if warehouse_item:
                if batch.batch_type == 'import':
                    warehouse_item.quantity += item.quantity
                else:
                    warehouse_item.quantity -= item.quantity
                warehouse_item.save()
            batch.calculate_total()
            batch.save()
            messages.success(request, 'Sản phẩm đã được thêm vào phiếu!')
            return redirect('dashboard:warehouse_batch_add_items', pk=batch.pk)
    else:
        form = WarehouseBatchItemForm()
    batch.calculate_total()
    return render(request, 'dashboard/warehouse/batch_add_items.html', {
        'batch': batch, 'form': form,
        'page_title': f'Thêm Sản Phẩm - Phiếu {batch.batch_number}',
    })

@admin_required
def warehouse_batch_detail(request, pk):
    batch = get_object_or_404(WarehouseBatch, pk=pk)
    batch.calculate_total()
    return render(request, 'dashboard/warehouse/batch_detail.html', {
        'batch': batch, 'page_title': f'Phiếu {batch.batch_number}',
    })

@admin_required
def warehouse_batch_print(request, pk):
    from .utils import export_batch_to_pdf
    batch = get_object_or_404(WarehouseBatch, pk=pk)
    batch.calculate_total()
    batch.is_printed = True
    batch.printed_at = timezone.now()
    batch.save()
    return export_batch_to_pdf(batch)

@admin_required
def warehouse_batch_export_excel(request, pk):
    from .utils import export_batch_to_excel
    batch = get_object_or_404(WarehouseBatch, pk=pk)
    batch.calculate_total()
    wb = export_batch_to_excel(batch)
    if wb is None:
        messages.error(request, 'openpyxl chưa được cài đặt.')
        return redirect('dashboard:warehouse_batch_detail', pk=batch.pk)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="phieu_{batch.batch_number}.xlsx"'
    wb.save(response)
    return response

@admin_required
def warehouse_batch_template_excel(request, warehouse_id):
    from .utils import generate_excel_template
    warehouse = get_object_or_404(Warehouse, pk=warehouse_id)
    wb = generate_excel_template(warehouse)
    if wb is None:
        messages.error(request, 'Lỗi: openpyxl chưa được cài đặt hoặc lỗi tạo template')
        return redirect('dashboard:warehouse_batch_create')
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="template_import_{warehouse.store.name}.xlsx"'
    wb.save(response)
    return response

@admin_required
@require_http_methods(["POST"])
def warehouse_batch_delete(request, pk):
    batch = get_object_or_404(WarehouseBatch, pk=pk)
    batch.delete()
    messages.success(request, 'Phiếu đã được xóa!')
    return redirect('dashboard:warehouse_batch_list')

@admin_required
def warehouse_batch_item_delete(request, pk):
    item = get_object_or_404(WarehouseBatchItem, pk=pk)
    batch = item.batch
    warehouse_item = WarehouseItem.objects.filter(
        warehouse=batch.warehouse, product=item.product
    ).first()
    if warehouse_item:
        if batch.batch_type == 'import':
            warehouse_item.quantity -= item.quantity
        else:
            warehouse_item.quantity += item.quantity
        warehouse_item.save()
    item.delete()
    batch.calculate_total()
    batch.save()
    messages.success(request, 'Sản phẩm đã được xóa khỏi phiếu!')
    return redirect('dashboard:warehouse_batch_add_items', pk=batch.pk)

@admin_required
def warehouse_import_excel(request):
    from .utils import import_warehouse_from_excel
    if request.method == 'POST':
        form = ImportExcelForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = request.FILES.get('excel_file')
            warehouse = form.cleaned_data['warehouse']
            supplier = form.cleaned_data.get('supplier', '')
            result = import_warehouse_from_excel(excel_file, warehouse, supplier)
            if result['success']:
                messages.success(request, f"Đã nhập {result['items_added']} sản phẩm. Phiếu: {result['batch_number']}")
                if result['errors']:
                    for error in result['errors']:
                        messages.warning(request, error)
                return redirect('dashboard:warehouse_batch_detail', pk=result['batch_id'])
            else:
                messages.error(request, f"Lỗi: {result['error']}")
    else:
        form = ImportExcelForm()
    return render(request, 'dashboard/warehouse/import_excel.html', {
        'form': form, 'page_title': 'Nhập Kho Từ Excel',
    })


# ==================== ABOUT MANAGEMENT (from dashboard1) ====================

@admin_required
def about_list(request):
    articles = About.objects.all().order_by('order', '-created_at')
    status_filter = request.GET.get('status')
    if status_filter:
        articles = articles.filter(status=status_filter)
    page = request.GET.get('page', 1)
    articles, paginator = paginate_queryset(articles, page)
    return render(request, 'dashboard/about/list.html', {
        'articles': articles, 'page_title': 'Quản Lý Trang Giới Thiệu', 'status_filter': status_filter,
    })

@admin_required
def about_create(request):
    if request.method == 'POST':
        form = AboutForm(request.POST, request.FILES)
        if form.is_valid():
            article = form.save(commit=False)
            article.author = request.user
            article.save()
            messages.success(request, 'Bài viết giới thiệu đã được tạo thành công!')
            return redirect('dashboard:about_list')
    else:
        form = AboutForm()
    return render(request, 'dashboard/about/form.html', {
        'form': form, 'page_title': 'Tạo Bài Viết Giới Thiệu Mới', 'action': 'create',
    })

@admin_required
def about_edit(request, pk):
    article = get_object_or_404(About, pk=pk)
    if request.method == 'POST':
        form = AboutForm(request.POST, request.FILES, instance=article)
        if form.is_valid():
            form.save()
            messages.success(request, 'Bài viết giới thiệu đã được cập nhật thành công!')
            return redirect('dashboard:about_list')
    else:
        form = AboutForm(instance=article)
    return render(request, 'dashboard/about/form.html', {
        'form': form, 'article': article, 'page_title': f'Chỉnh Sửa: {article.title}', 'action': 'edit',
    })

@admin_required
def about_delete(request, pk):
    article = get_object_or_404(About, pk=pk)
    if request.method == 'POST':
        article.delete()
        messages.success(request, 'Bài viết giới thiệu đã được xóa thành công!')
        return redirect('dashboard:about_list')
    return render(request, 'dashboard/about/delete.html', {
        'article': article, 'page_title': 'Xóa Bài Viết Giới Thiệu',
    })

@admin_required
def about_detail(request, pk):
    article = get_object_or_404(About, pk=pk)
    return render(request, 'dashboard/about/detail.html', {
        'article': article, 'page_title': f'Chi Tiết: {article.title}',
    })

@admin_required
def about_import(request):
    if request.method == 'POST':
        form = AboutImportForm(request.POST, request.FILES)
        if form.is_valid():
            import_type = form.cleaned_data['import_type']
            if import_type == 'word':
                word_file = form.cleaned_data['word_file']
                try:
                    article = About.objects.create(
                        title=f"Import from {word_file.name}",
                        slug=f"import-{word_file.name}",
                        content="Content imported from Word file",
                        source_type='word', author=request.user, status='draft'
                    )
                    messages.success(request, 'File Word đã được nhập thành công!')
                    return redirect('dashboard:about_edit', pk=article.pk)
                except Exception as e:
                    messages.error(request, f'Lỗi khi đọc file Word: {str(e)}')
            elif import_type == 'external':
                external_url = form.cleaned_data['external_url']
                try:
                    article = About.objects.create(
                        title=f"Import from {external_url}",
                        slug=f"import-{external_url.replace('https://', '').replace('/', '-')}",
                        content=f"Content imported from: {external_url}",
                        external_link=external_url, source_type='external',
                        author=request.user, status='draft'
                    )
                    messages.success(request, 'Nội dung từ URL đã được nhập thành công!')
                    return redirect('dashboard:about_edit', pk=article.pk)
                except Exception as e:
                    messages.error(request, f'Lỗi khi lấy nội dung từ URL: {str(e)}')
    else:
        form = AboutImportForm()
    return render(request, 'dashboard/about/import.html', {
        'form': form, 'page_title': 'Nhập Bài Viết Từ Nguồn Bên Ngoài',
    })


# ==================== CUSTOMER PROFILE (from dashboard1) ====================

@login_required
def customer_profile_view(request):
    profile, created = CustomerProfile.objects.get_or_create(user=request.user)
    return render(request, 'customer/profile.html', {
        'profile': profile, 'page_title': 'Thông Tin Cá Nhân',
    })

@login_required
def customer_profile_edit(request):
    profile, created = CustomerProfile.objects.get_or_create(user=request.user)
    if request.method == 'POST':
        profile_form = CustomerProfileForm(request.POST, request.FILES, instance=profile)
        user_form = UserBasicInfoForm(request.POST, instance=request.user)
        if profile_form.is_valid() and user_form.is_valid():
            profile_form.save()
            user_form.save()
            messages.success(request, 'Thông tin cá nhân đã được cập nhật thành công!')
            return redirect('customer_profile')
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        profile_form = CustomerProfileForm(instance=profile)
        user_form = UserBasicInfoForm(instance=request.user)
    return render(request, 'customer/profile_edit_enhanced.html', {
        'profile_form': profile_form, 'user_form': user_form,
        'profile': profile, 'page_title': 'Chỉnh Sửa Thông Tin Cá Nhân',
    })

@admin_required
def customer_profile_list(request):
    profiles = CustomerProfile.objects.select_related('user').order_by('-created_at')
    search_query = request.GET.get('search', '')
    if search_query:
        profiles = profiles.filter(
            Q(user__username__icontains=search_query) |
            Q(user__email__icontains=search_query) |
            Q(user__first_name__icontains=search_query) |
            Q(user__last_name__icontains=search_query) |
            Q(phone__icontains=search_query) |
            Q(address__icontains=search_query)
        )
    verified_filter = request.GET.get('verified', '')
    if verified_filter == 'verified':
        profiles = profiles.filter(is_verified=True)
    elif verified_filter == 'unverified':
        profiles = profiles.filter(is_verified=False)
    status_filter = request.GET.get('status', '')
    if status_filter == 'active':
        profiles = profiles.filter(user__is_active=True)
    elif status_filter == 'inactive':
        profiles = profiles.filter(user__is_active=False)
    gender_filter = request.GET.get('gender', '')
    if gender_filter:
        profiles = profiles.filter(gender=gender_filter)
    vip_filter = request.GET.get('vip', '')
    if vip_filter == 'vip':
        profiles = profiles.filter(is_premium=True)
    elif vip_filter == 'normal':
        profiles = profiles.filter(is_premium=False)

    total_customers = CustomerProfile.objects.count()
    verified_count = CustomerProfile.objects.filter(is_verified=True).count()
    vip_count = CustomerProfile.objects.filter(is_premium=True).count()
    active_count = CustomerProfile.objects.filter(user__is_active=True).count()
    new_this_month = CustomerProfile.objects.filter(
        created_at__month=timezone.now().month, created_at__year=timezone.now().year
    ).count()
    verification_rate = round((verified_count / total_customers) * 100, 1) if total_customers > 0 else 0

    page = request.GET.get('page', 1)
    profiles, paginator = paginate_queryset(profiles, page, 20)
    return render(request, 'dashboard/customer_profiles/list.html', {
        'profiles': profiles, 'page_title': 'Quản Lý Thông Tin Người Dùng',
        'search_query': search_query, 'verified_filter': verified_filter,
        'status_filter': status_filter, 'gender_filter': gender_filter,
        'vip_filter': vip_filter,
        'stats': {
            'total': total_customers, 'verified': verified_count,
            'vip': vip_count, 'active': active_count,
            'new_this_month': new_this_month, 'verification_rate': verification_rate,
        },
        'gender_choices': CustomerProfile._meta.get_field('gender').choices,
    })

@admin_required
def customer_profile_detail(request, pk):
    profile = get_object_or_404(CustomerProfile, pk=pk)
    customer_orders = Order.objects.filter(user=profile.user).order_by('-created_at')[:10]
    total_orders = Order.objects.filter(user=profile.user).count()
    total_spent = Order.objects.filter(user=profile.user).aggregate(total=Sum('total_amount'))['total'] or 0
    orders_this_month = Order.objects.filter(
        user=profile.user, created_at__month=timezone.now().month,
        created_at__year=timezone.now().year
    ).count()
    recent_orders_with_items = []
    for order in customer_orders[:5]:
        items = OrderItem.objects.filter(order=order).select_related('product')
        recent_orders_with_items.append({'order': order, 'items': items, 'item_count': items.count()})
    return render(request, 'dashboard/customer_profiles/detail.html', {
        'profile': profile, 'page_title': f'Chi Tiết: {profile.display_name}',
        'customer_orders': customer_orders, 'recent_orders_with_items': recent_orders_with_items,
        'order_stats': {
            'total_orders': total_orders, 'total_spent': total_spent,
            'orders_this_month': orders_this_month,
        },
        'loyalty_points': profile.loyalty_points,
        'profile_completion': profile.get_completion_percentage(),
    })

@admin_required
def customer_profile_edit(request, pk):
    profile = get_object_or_404(CustomerProfile, pk=pk)
    if request.method == 'POST':
        profile_form = CustomerProfileForm(request.POST, request.FILES, instance=profile)
        user_form = UserBasicInfoForm(request.POST, instance=profile.user)
        if profile_form.is_valid() and user_form.is_valid():
            profile.user.is_active = request.POST.get('is_active') == 'on'
            profile.user.is_staff = request.POST.get('is_staff') == 'on'
            profile.is_verified = request.POST.get('is_verified') == 'on'
            profile_form.save()
            user_form.save()
            messages.success(request, f'Thông tin người dùng {profile.display_name} đã được cập nhật thành công!')
            return redirect('dashboard:customer_profile_detail', pk=profile.pk)
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        profile_form = CustomerProfileForm(instance=profile)
        user_form = UserBasicInfoForm(instance=profile.user)
    return render(request, 'dashboard/customer_profiles/edit.html', {
        'profile_form': profile_form, 'user_form': user_form,
        'profile': profile, 'page_title': f'Chỉnh Sửa: {profile.display_name}',
    })

@admin_required
def customer_profile_delete(request, pk):
    profile = get_object_or_404(CustomerProfile, pk=pk)
    user = profile.user
    if request.method == 'POST':
        username = profile.display_name
        profile.delete()
        user.delete()
        messages.success(request, f'Người dùng {username} đã được xóa thành công!')
        return redirect('dashboard:customer_profile_list')
    return render(request, 'dashboard/customer_profiles/delete.html', {
        'profile': profile, 'page_title': f'Xóa: {profile.display_name}',
    })

@admin_required
def customer_profile_toggle_verification(request, pk):
    profile = get_object_or_404(CustomerProfile, pk=pk)
    profile.is_verified = not profile.is_verified
    profile.save()
    v_status = "xác thực" if profile.is_verified else "hủy xác thực"
    messages.success(request, f'Đã {v_status} tài khoản của {profile.display_name}!')
    return redirect('dashboard:customer_profile_list')


# ==================== UNIFIED CUSTOMERS ====================

@admin_required
def customers_unified_list(request):
    """Unified customer list combining User + CustomerProfile data"""
    users = User.objects.select_related('profile').all()

    # Stats
    total_users = User.objects.count()
    active_count = User.objects.filter(is_active=True).count()
    verified_count = CustomerProfile.objects.filter(is_verified=True).count()
    vip_count = CustomerProfile.objects.filter(is_premium=True).count()
    new_this_month = User.objects.filter(
        date_joined__month=timezone.now().month,
        date_joined__year=timezone.now().year
    ).count()

    # Filters
    search_query = request.GET.get('search', '')
    if search_query:
        users = users.filter(
            Q(email__icontains=search_query) |
            Q(first_name__icontains=search_query) |
            Q(last_name__icontains=search_query) |
            Q(phone__icontains=search_query)
        )
    status_filter = request.GET.get('status', '')
    if status_filter == 'active':
        users = users.filter(is_active=True)
    elif status_filter == 'inactive':
        users = users.filter(is_active=False)
    role_filter = request.GET.get('role', '')
    if role_filter == 'admin':
        users = users.filter(is_staff=True)
    elif role_filter == 'user':
        users = users.filter(is_staff=False)
    vip_filter = request.GET.get('vip', '')
    if vip_filter == 'vip':
        users = users.filter(profile__is_premium=True)
    elif vip_filter == 'normal':
        users = users.filter(profile__is_premium=False)

    page = request.GET.get('page', 1)
    page_obj, paginator = paginate_queryset(users.order_by('-date_joined'), page, 15)

    return render(request, 'dashboard/customers_unified/list.html', {
        'page_obj': page_obj, 'paginator': paginator,
        'page_title': 'Quản Lý Khách Hàng',
        'search_query': search_query,
        'status_filter': status_filter,
        'role_filter': role_filter,
        'vip_filter': vip_filter,
        'stats': {
            'total': total_users, 'active': active_count,
            'verified': verified_count, 'vip': vip_count,
            'new_this_month': new_this_month,
        },
    })


@admin_required
def customers_unified_detail(request, pk):
    """Unified customer detail combining user + profile + orders"""
    user_obj = get_object_or_404(User.objects.select_related('profile'), pk=pk)
    profile = user_obj.profile if hasattr(user_obj, 'profile') and user_obj.profile else None

    # Order stats
    total_orders = Order.objects.filter(user=user_obj).count()
    total_spent = Order.objects.filter(user=user_obj).aggregate(total=Sum('total_amount'))['total'] or 0
    orders_this_month = Order.objects.filter(
        user=user_obj, created_at__month=timezone.now().month,
        created_at__year=timezone.now().year
    ).count()

    # Recent orders with items
    customer_orders = Order.objects.filter(user=user_obj).order_by('-created_at')[:10]
    recent_orders_with_items = []
    for order in customer_orders[:5]:
        items = OrderItem.objects.filter(order=order).select_related('product')
        recent_orders_with_items.append({'order': order, 'items': items, 'item_count': items.count()})

    profile_completion = profile.get_completion_percentage() if profile else 0
    loyalty_points = profile.loyalty_points if profile else 0

    return render(request, 'dashboard/customers_unified/detail.html', {
        'user_obj': user_obj, 'profile': profile,
        'customer_orders': customer_orders,
        'recent_orders_with_items': recent_orders_with_items,
        'page_title': f'Khách Hàng: {user_obj.get_full_name() or user_obj.email}',
        'order_stats': {
            'total_orders': total_orders, 'total_spent': total_spent,
            'orders_this_month': orders_this_month,
        },
        'loyalty_points': loyalty_points,
        'profile_completion': profile_completion,
    })


@admin_required
def customers_unified_create(request):
    """Create new user with profile"""
    if request.method == 'POST':
        user_form = UserForm(request.POST)
        profile_form = CustomerProfileForm(request.POST, request.FILES)
        if user_form.is_valid() and profile_form.is_valid():
            user = user_form.save(commit=False)
            password = user_form.cleaned_data.get('password')
            if password:
                user.set_password(password)
            else:
                user.set_password('defaultpass123')
            user.save()
            profile = profile_form.save(commit=False)
            profile.user = user
            profile.save()
            messages.success(request, f'Khách hàng "{user.email}" đã được tạo thành công!')
            return redirect('dashboard:customers_unified_detail', pk=user.pk)
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        user_form = UserForm()
        profile_form = CustomerProfileForm()
    return render(request, 'dashboard/customers_unified/form.html', {
        'user_form': user_form, 'profile_form': profile_form,
        'page_title': 'Thêm Khách Hàng', 'is_create': True,
    })


@admin_required
def customers_unified_edit(request, pk):
    """Edit user + profile together"""
    user_obj = get_object_or_404(User, pk=pk)
    profile, created = CustomerProfile.objects.get_or_create(user=user_obj)

    if request.method == 'POST':
        user_form = UserForm(request.POST, instance=user_obj)
        profile_form = CustomerProfileForm(request.POST, request.FILES, instance=profile)
        if user_form.is_valid() and profile_form.is_valid():
            u = user_form.save(commit=False)
            password = user_form.cleaned_data.get('password')
            if password:
                u.set_password(password)
            u.is_active = request.POST.get('is_active') == 'on'
            u.is_staff = request.POST.get('is_staff') == 'on'
            u.save()
            profile.is_verified = request.POST.get('is_verified') == 'on'
            profile.is_premium = request.POST.get('is_premium') == 'on'
            profile_form.save()
            messages.success(request, f'Khách hàng đã được cập nhật!')
            return redirect('dashboard:customers_unified_detail', pk=user_obj.pk)
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        user_form = UserForm(instance=user_obj)
        profile_form = CustomerProfileForm(instance=profile)
    return render(request, 'dashboard/customers_unified/form.html', {
        'user_form': user_form, 'profile_form': profile_form,
        'user_obj': user_obj, 'profile': profile,
        'page_title': f'Chỉnh Sửa: {user_obj.get_full_name() or user_obj.email}',
        'is_create': False,
    })


@admin_required
@require_http_methods(["POST"])
def customers_unified_delete(request, pk):
    """Delete user and their profile"""
    user_obj = get_object_or_404(User, pk=pk)
    name = user_obj.get_full_name() or user_obj.email
    user_obj.delete()
    messages.success(request, f'Khách hàng "{name}" đã được xóa!')
    return redirect('dashboard:customers_unified_list')


@admin_required
def customers_unified_toggle_verification(request, pk):
    """Toggle user verification status"""
    user_obj = get_object_or_404(User, pk=pk)
    profile, created = CustomerProfile.objects.get_or_create(user=user_obj)
    profile.is_verified = not profile.is_verified
    profile.save()
    v_status = "xác thực" if profile.is_verified else "hủy xác thực"
    messages.success(request, f'Đã {v_status} tài khoản của {user_obj.get_full_name() or user_obj.email}!')
    return redirect('dashboard:customers_unified_list')


# ==================== NEWS MANAGEMENT (from dashboard1) ====================

@admin_required
def news_list(request):
    news_items = News.objects.all().order_by('-created_at')
    status_filter = request.GET.get('status')
    if status_filter:
        news_items = news_items.filter(status=status_filter)
    category_filter = request.GET.get('category')
    if category_filter:
        news_items = news_items.filter(category=category_filter)
    search_query = request.GET.get('search')
    if search_query:
        news_items = news_items.filter(title__icontains=search_query)
    return render(request, 'dashboard/news/list.html', {
        'news_items': news_items, 'page_title': 'Quản Lý Tin Tức',
        'status_choices': News.STATUS_CHOICES, 'category_choices': News.CATEGORY_CHOICES,
        'current_status': status_filter, 'current_category': category_filter,
        'search_query': search_query,
    })

@admin_required
def news_create(request):
    if request.method == 'POST':
        form = NewsForm(request.POST, request.FILES)
        if form.is_valid():
            news = form.save(commit=False)
            news.author = request.user
            if news.status == 'published' and not news.published_at:
                news.published_at = timezone.now()
            news.save()
            messages.success(request, 'Tin tức đã được tạo thành công!')
            return redirect('dashboard:news_list')
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        form = NewsForm()
    return render(request, 'dashboard/news/form.html', {
        'form': form, 'page_title': 'Thêm Tin Tức Mới',
    })

@admin_required
def news_edit(request, pk):
    news = get_object_or_404(News, pk=pk)
    if request.method == 'POST':
        form = NewsForm(request.POST, request.FILES, instance=news)
        if form.is_valid():
            news = form.save(commit=False)
            if news.status == 'published' and not news.published_at:
                news.published_at = timezone.now()
            news.save()
            messages.success(request, 'Tin tức đã được cập nhật thành công!')
            return redirect('dashboard:news_list')
        else:
            messages.error(request, 'Vui lòng sửa các lỗi bên dưới.')
    else:
        form = NewsForm(instance=news)
    return render(request, 'dashboard/news/form.html', {
        'form': form, 'news': news, 'page_title': f'Sửa: {news.title}',
    })

@admin_required
def news_delete(request, pk):
    news = get_object_or_404(News, pk=pk)
    if request.method == 'POST':
        title = news.title
        news.delete()
        messages.success(request, f'Tin tức "{title}" đã được xóa thành công!')
        return redirect('dashboard:news_list')
    return render(request, 'dashboard/news/delete.html', {
        'news': news, 'page_title': f'Xóa: {news.title}',
    })

@admin_required
def news_toggle_status(request, pk):
    news = get_object_or_404(News, pk=pk)
    if news.status == 'published':
        news.status = 'draft'
        messages.success(request, f'Tin tức "{news.title}" đã chuyển sang bản nháp!')
    else:
        news.status = 'published'
        if not news.published_at:
            news.published_at = timezone.now()
        messages.success(request, f'Tin tức "{news.title}" đã được xuất bản!')
    news.save()
    return redirect('dashboard:news_list')

@admin_required
def news_toggle_featured(request, pk):
    news = get_object_or_404(News, pk=pk)
    news.is_featured = not news.is_featured
    news.save()
    f_status = "nổi bật" if news.is_featured else "bình thường"
    messages.success(request, f'Tin tức "{news.title}" đã chuyển sang trạng thái {f_status}!')
    return redirect('dashboard:news_list')


# ==================== REVIEW MANAGEMENT (from dashboard1) ====================

@admin_required
def review_list(request):
    reviews = Review.objects.select_related('user', 'order_item__product', 'order_item__order').prefetch_related('images').all()
    status_filter = request.GET.get('status')
    if status_filter:
        reviews = reviews.filter(status=status_filter)
    rating_filter = request.GET.get('rating')
    if rating_filter:
        reviews = reviews.filter(rating=rating_filter)
    search_query = request.GET.get('q')
    if search_query:
        reviews = reviews.filter(
            Q(content__icontains=search_query) |
            Q(user__username__icontains=search_query) |
            Q(order_item__product__name__icontains=search_query)
        )
    stats = {
        'total': Review.objects.count(),
        'pending': Review.objects.filter(status='pending').count(),
        'approved': Review.objects.filter(status='approved').count(),
        'rejected': Review.objects.filter(status='rejected').count(),
        'avg_rating': Review.objects.filter(status='approved').aggregate(avg=Avg('rating'))['avg'] or 0,
    }
    paginator = Paginator(reviews, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    return render(request, 'dashboard/reviews/list.html', {
        'reviews': page_obj, 'page_obj': page_obj, 'stats': stats, 'page_title': 'Quản lý Đánh giá',
        'status_filter': status_filter, 'rating_filter': rating_filter, 'search_query': search_query,
    })

@admin_required
def review_detail(request, pk):
    review = get_object_or_404(Review.objects.select_related('user', 'order_item__product', 'order_item__order'), pk=pk)
    if request.method == 'POST':
        reply_form = ReviewReplyForm(request.POST)
        if reply_form.is_valid():
            review.admin_reply = reply_form.cleaned_data['admin_reply']
            review.admin_reply_at = timezone.now()
            review.save()
            messages.success(request, 'Đã gửi phản hồi!')
            return redirect('dashboard:review_detail', pk=pk)
    else:
        reply_form = ReviewReplyForm(initial={'admin_reply': review.admin_reply})
    return render(request, 'dashboard/reviews/detail.html', {
        'review': review, 'reply_form': reply_form, 'page_title': f'Đánh giá #{review.id}',
    })

@admin_required
def review_approve(request, pk):
    if request.method != 'POST':
        return redirect('dashboard:review_list')
    review = get_object_or_404(Review, pk=pk)
    review.status = 'approved'
    review.save()
    messages.success(request, f'Đã duyệt đánh giá #{review.id}!')
    return redirect('dashboard:review_list')

@admin_required
def review_reject(request, pk):
    if request.method != 'POST':
        return redirect('dashboard:review_list')
    review = get_object_or_404(Review, pk=pk)
    review.status = 'rejected'
    review.save()
    messages.success(request, f'Đã từ chối đánh giá #{review.id}!')
    return redirect('dashboard:review_list')

@admin_required
def review_delete(request, pk):
    review = get_object_or_404(Review, pk=pk)
    if request.method == 'POST':
        review.delete()
        messages.success(request, 'Đã xóa đánh giá!')
        return redirect('dashboard:review_list')
    return render(request, 'dashboard/reviews/delete.html', {
        'review': review, 'page_title': 'Xóa đánh giá',
    })


# ==================== ORDER REVIEW MANAGEMENT (from dashboard1) ====================

@admin_required
def order_review_list(request):
    order_reviews = OrderReview.objects.select_related('order', 'user').all().order_by('-created_at')
    status_filter = request.GET.get('status')
    if status_filter:
        order_reviews = order_reviews.filter(status=status_filter)
    rating_filter = request.GET.get('rating')
    if rating_filter:
        order_reviews = order_reviews.filter(overall_rating=rating_filter)
    search_query = request.GET.get('q')
    if search_query:
        order_reviews = order_reviews.filter(
            Q(order__id__icontains=search_query) |
            Q(user__username__icontains=search_query) |
            Q(content__icontains=search_query)
        )
    page_obj, paginator = paginate_queryset(order_reviews, request.GET.get('page'), 10)
    stats = {
        'total': OrderReview.objects.count(),
        'pending': OrderReview.objects.filter(status='pending').count(),
        'approved': OrderReview.objects.filter(status='approved').count(),
        'rejected': OrderReview.objects.filter(status='rejected').count(),
        'avg_rating': OrderReview.objects.filter(status='approved').aggregate(avg=Avg('overall_rating'))['avg'] or 0,
    }
    return render(request, 'dashboard/order_reviews/list.html', {
        'order_reviews': page_obj, 'paginator': paginator, 'stats': stats,
        'page_title': 'Quản lý Đánh giá Đơn hàng',
        'status_filter': status_filter, 'rating_filter': rating_filter, 'search_query': search_query,
    })

@admin_required
def order_review_detail(request, pk):
    order_review = get_object_or_404(OrderReview.objects.select_related('order', 'user'), pk=pk)
    if request.method == 'POST':
        reply_content = request.POST.get('admin_reply')
        if reply_content:
            order_review.admin_reply = reply_content
            order_review.admin_reply_at = timezone.now()
            order_review.save()
            messages.success(request, 'Đã gửi phản hồi!')
            return redirect('dashboard:order_review_detail', pk=pk)
    return render(request, 'dashboard/order_reviews/detail.html', {
        'order_review': order_review, 'page_title': f'Đánh giá Đơn hàng #{order_review.order.id}',
    })

@admin_required
def order_review_approve(request, pk):
    if request.method != 'POST':
        return redirect('dashboard:order_review_list')
    order_review = get_object_or_404(OrderReview, pk=pk)
    order_review.status = 'approved'
    order_review.save()
    messages.success(request, f'Đã duyệt đánh giá đơn hàng #{order_review.order.id}!')
    return redirect('dashboard:order_review_list')

@admin_required
def order_review_reject(request, pk):
    if request.method != 'POST':
        return redirect('dashboard:order_review_list')
    order_review = get_object_or_404(OrderReview, pk=pk)
    order_review.status = 'rejected'
    order_review.save()
    messages.success(request, f'Đã từ chối đánh giá đơn hàng #{order_review.order.id}!')
    return redirect('dashboard:order_review_list')

@admin_required
def order_review_delete(request, pk):
    order_review = get_object_or_404(OrderReview, pk=pk)
    if request.method == 'POST':
        order_review.delete()
        messages.success(request, 'Đã xóa đánh giá đơn hàng!')
        return redirect('dashboard:order_review_list')
    return render(request, 'dashboard/order_reviews/delete.html', {
        'order_review': order_review, 'page_title': 'Xóa đánh giá đơn hàng',
    })


# ==================== CONTACT SETTINGS ====================

@admin_required
def contact_settings(request):
    obj = ContactSettings.objects.first()
    if not obj:
        obj = ContactSettings.objects.create()

    if request.method == 'POST':
        form = ContactSettingsForm(request.POST, instance=obj)
        if form.is_valid():
            form.save()
            messages.success(request, 'Đã cập nhật cài đặt Liên hệ!')
            return redirect('dashboard:contact_settings')
    else:
        form = ContactSettingsForm(instance=obj)

    return render(request, 'dashboard/contact_settings.html', {
        'form': form,
        'settings': obj,
        'page_title': 'Cài đặt Liên hệ',
    })


# ==================== SUPPLIER MANAGEMENT ====================

@admin_required
def supplier_list(request):
    """Danh sách nhà cung cấp"""
    query = request.GET.get('q', '')
    suppliers = Supplier.objects.all()
    
    if query:
        suppliers = suppliers.filter(
            Q(name__icontains=query) |
            Q(contact_person__icontains=query) |
            Q(email__icontains=query) |
            Q(phone__icontains=query)
        )
    
    suppliers = suppliers.order_by('name')
    page_obj, paginator = paginate_queryset(suppliers, request.GET.get('page'), 20)
    
    return render(request, 'dashboard/suppliers/list.html', {
        'page_obj': page_obj,
        'query': query,
        'page_title': 'Nhà cung cấp',
    })

@admin_required
def supplier_create(request):
    """Tạo nhà cung cấp mới"""
    if request.method == 'POST':
        form = SupplierForm(request.POST)
        if form.is_valid():
            supplier = form.save()
            messages.success(request, f'Đã tạo nhà cung cấp {supplier.name}!')
            return redirect('dashboard:supplier_list')
    else:
        form = SupplierForm()
    
    return render(request, 'dashboard/suppliers/form.html', {
        'form': form,
        'page_title': 'Thêm nhà cung cấp',
    })

@admin_required
def supplier_edit(request, pk):
    """Chỉnh sửa nhà cung cấp"""
    supplier = get_object_or_404(Supplier, pk=pk)
    
    if request.method == 'POST':
        form = SupplierForm(request.POST, instance=supplier)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật nhà cung cấp {supplier.name}!')
            return redirect('dashboard:supplier_list')
    else:
        form = SupplierForm(instance=supplier)
    
    return render(request, 'dashboard/suppliers/form.html', {
        'form': form,
        'supplier': supplier,
        'page_title': 'Chỉnh sửa nhà cung cấp',
    })

@admin_required
def supplier_detail(request, pk):
    """Chi tiết nhà cung cấp"""
    supplier = get_object_or_404(Supplier, pk=pk)
    
    # Lấy các đơn mua hàng gần đây
    purchase_orders = PurchaseOrder.objects.filter(
        supplier=supplier
    ).order_by('-order_date')[:10]
    
    return render(request, 'dashboard/suppliers/detail.html', {
        'supplier': supplier,
        'purchase_orders': purchase_orders,
        'page_title': f'Chi tiết nhà cung cấp: {supplier.name}',
    })

@admin_required
def supplier_delete(request, pk):
    """Xóa nhà cung cấp"""
    supplier = get_object_or_404(Supplier, pk=pk)
    
    if request.method == 'POST':
        # Kiểm tra xem có đơn mua hàng nào không
        if PurchaseOrder.objects.filter(supplier=supplier).exists():
            messages.error(request, 'Không thể xóa nhà cung cấp đã có đơn mua hàng!')
            return redirect('dashboard:supplier_detail', pk=pk)
        if GoodsReceipt.objects.filter(supplier=supplier).exists():
            messages.error(request, 'Không thể xóa nhà cung cấp đã có phiếu nhập hàng!')
            return redirect('dashboard:supplier_detail', pk=pk)
        
        supplier.delete()
        messages.success(request, f'Đã xóa nhà cung cấp {supplier.name}!')
        return redirect('dashboard:supplier_list')
    
    return render(request, 'dashboard/suppliers/delete.html', {
        'supplier': supplier,
        'page_title': 'Xóa nhà cung cấp',
    })


# ==================== PURCHASE ORDER MANAGEMENT ====================

@admin_required
def purchase_order_list(request):
    """Danh sách đơn mua hàng"""
    query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    supplier_filter = request.GET.get('supplier', '')
    
    purchase_orders = PurchaseOrder.objects.select_related('supplier').prefetch_related('items')
    
    if query:
        purchase_orders = purchase_orders.filter(
            Q(po_number__icontains=query) |
            Q(supplier__name__icontains=query) |
            Q(note__icontains=query)
        )
    
    if status_filter:
        purchase_orders = purchase_orders.filter(status=status_filter)
    
    if supplier_filter:
        purchase_orders = purchase_orders.filter(supplier_id=supplier_filter)
    
    purchase_orders = purchase_orders.order_by('-order_date')
    page_obj, paginator = paginate_queryset(purchase_orders, request.GET.get('page'), 20)
    
    suppliers = Supplier.objects.filter(is_active=True)
    
    return render(request, 'dashboard/purchase_orders/list.html', {
        'page_obj': page_obj,
        'query': query,
        'suppliers': suppliers,
        'status_filter': status_filter,
        'supplier_filter': supplier_filter,
        'page_title': 'Đơn mua hàng',
    })

@admin_required
def purchase_order_create(request):
    """Tạo đơn mua hàng mới"""
    if request.method == 'POST':
        form = PurchaseOrderForm(request.POST)
        if form.is_valid():
            purchase_order = form.save()
            messages.success(request, f'Đã tạo đơn mua hàng {purchase_order.po_number}!')
            return redirect('dashboard:purchase_order_detail', pk=purchase_order.pk)
    else:
        initial = {}
        sup = request.GET.get('supplier')
        if sup and sup.isdigit():
            initial['supplier'] = int(sup)
        form = PurchaseOrderForm(initial=initial)
    
    return render(request, 'dashboard/purchase_orders/form.html', {
        'form': form,
        'page_title': 'Thêm đơn mua hàng',
    })

@admin_required
def purchase_order_edit(request, pk):
    """Chỉnh sửa đơn mua hàng"""
    purchase_order = get_object_or_404(PurchaseOrder, pk=pk)
    
    # Chỉ cho phép sửa đơn ở trạng thái nháp hoặc chờ duyệt
    if purchase_order.status not in ('draft', 'pending'):
        messages.error(request, 'Chỉ có thể sửa đơn ở trạng thái Nháp hoặc Chờ duyệt!')
        return redirect('dashboard:purchase_order_detail', pk=pk)
    
    if request.method == 'POST':
        form = PurchaseOrderForm(request.POST, instance=purchase_order)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật đơn mua hàng {purchase_order.po_number}!')
            return redirect('dashboard:purchase_order_detail', pk=pk)
    else:
        form = PurchaseOrderForm(instance=purchase_order)
    
    return render(request, 'dashboard/purchase_orders/form.html', {
        'form': form,
        'purchase_order': purchase_order,
        'page_title': 'Chỉnh sửa đơn mua hàng',
    })

@admin_required
def purchase_order_detail(request, pk):
    """Chi tiết đơn mua hàng"""
    purchase_order = get_object_or_404(
        PurchaseOrder.objects.select_related('supplier', 'store'),
        pk=pk
    )
    items = purchase_order.items.select_related('product')
    
    return render(request, 'dashboard/purchase_orders/detail.html', {
        'purchase_order': purchase_order,
        'items': items,
        'page_title': f'Chi tiết đơn mua hàng: {purchase_order.po_number}',
    })

@admin_required
def purchase_order_confirm(request, pk):
    """Xác nhận đơn mua hàng"""
    purchase_order = get_object_or_404(PurchaseOrder, pk=pk)
    
    if purchase_order.status not in ('draft', 'pending'):
        messages.error(request, 'Chỉ có thể xác nhận đơn ở trạng thái Nháp hoặc Chờ duyệt!')
        return redirect('dashboard:purchase_order_detail', pk=pk)

    if request.method == 'POST':
        purchase_order.status = 'approved'
        purchase_order.save()
        messages.success(request, f'Đã duyệt đơn mua hàng {purchase_order.po_number}!')
        return redirect('dashboard:purchase_order_detail', pk=pk)
    
    return render(request, 'dashboard/purchase_orders/confirm.html', {
        'purchase_order': purchase_order,
        'page_title': 'Xác nhận đơn mua hàng',
    })

@admin_required
def purchase_order_cancel(request, pk):
    """Hủy đơn mua hàng"""
    purchase_order = get_object_or_404(PurchaseOrder, pk=pk)
    
    if purchase_order.status not in ('draft', 'pending', 'approved'):
        messages.error(request, 'Không thể hủy đơn hàng ở trạng thái này!')
        return redirect('dashboard:purchase_order_detail', pk=pk)

    if request.method == 'POST':
        purchase_order.status = 'cancelled'
        purchase_order.save()
        messages.success(request, f'Đã hủy đơn mua hàng {purchase_order.po_number}!')
        return redirect('dashboard:purchase_order_detail', pk=pk)
    
    return render(request, 'dashboard/purchase_orders/cancel.html', {
        'purchase_order': purchase_order,
        'page_title': 'Hủy đơn mua hàng',
    })

@admin_required
def purchase_order_delete(request, pk):
    """Xóa đơn mua hàng"""
    purchase_order = get_object_or_404(PurchaseOrder, pk=pk)
    
    if purchase_order.status not in ('draft', 'pending'):
        messages.error(request, 'Chỉ có thể xóa đơn ở trạng thái Nháp hoặc Chờ duyệt!')
        return redirect('dashboard:purchase_order_detail', pk=pk)

    if request.method == 'POST':
        purchase_order.delete()
        messages.success(request, 'Đã xóa đơn mua hàng!')
        return redirect('dashboard:purchase_order_list')
    
    return render(request, 'dashboard/purchase_orders/delete.html', {
        'purchase_order': purchase_order,
        'page_title': 'Xóa đơn mua hàng',
    })


# ==================== GOODS RECEIPT MANAGEMENT ====================

@admin_required
def goods_receipt_list(request):
    """Danh sách phiếu nhập hàng"""
    query = request.GET.get('q', '')
    status_filter = request.GET.get('status', '')
    purchase_order_filter = request.GET.get('purchase_order', '')
    
    goods_receipts = GoodsReceipt.objects.select_related(
        'purchase_order', 'purchase_order__supplier', 'supplier', 'warehouse', 'warehouse__store'
    ).prefetch_related('items')
    
    if query:
        goods_receipts = goods_receipts.filter(
            Q(receipt_number__icontains=query) |
            Q(purchase_order__po_number__icontains=query) |
            Q(purchase_order__supplier__name__icontains=query) |
            Q(supplier__name__icontains=query) |
            Q(note__icontains=query)
        )
    
    if status_filter:
        goods_receipts = goods_receipts.filter(status=status_filter)
    
    if purchase_order_filter:
        goods_receipts = goods_receipts.filter(purchase_order_id=purchase_order_filter)
    
    goods_receipts = goods_receipts.order_by('-receipt_date')
    page_obj, paginator = paginate_queryset(goods_receipts, request.GET.get('page'), 20)
    
    purchase_orders = PurchaseOrder.objects.filter(status__in=['approved', 'pending', 'draft']).order_by('-order_date')
    
    return render(request, 'dashboard/goods_receipts/list.html', {
        'page_obj': page_obj,
        'query': query,
        'purchase_orders': purchase_orders,
        'status_filter': status_filter,
        'purchase_order_filter': purchase_order_filter,
        'page_title': 'Phiếu nhập hàng',
    })

@admin_required
def goods_receipt_create(request):
    """Tạo phiếu nhập hàng mới"""
    if request.method == 'POST':
        form = GoodsReceiptForm(request.POST)
        if form.is_valid():
            goods_receipt = form.save()
            messages.success(request, f'Đã tạo phiếu nhập hàng {goods_receipt.receipt_number}!')
            return redirect('dashboard:goods_receipt_detail', pk=goods_receipt.pk)
    else:
        form = GoodsReceiptForm()
    
    return render(request, 'dashboard/goods_receipts/form.html', {
        'form': form,
        'page_title': 'Thêm phiếu nhập hàng',
    })

@admin_required
def goods_receipt_edit(request, pk):
    """Chỉnh sửa phiếu nhập hàng"""
    goods_receipt = get_object_or_404(GoodsReceipt, pk=pk)
    
    # Chỉ cho phép sửa phiếu ở trạng thái nháp
    if goods_receipt.status != 'draft':
        messages.error(request, 'Chỉ có thể sửa phiếu nhập ở trạng thái Nháp!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)
    
    if request.method == 'POST':
        form = GoodsReceiptForm(request.POST, instance=goods_receipt)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật phiếu nhập hàng {goods_receipt.receipt_number}!')
            return redirect('dashboard:goods_receipt_detail', pk=pk)
    else:
        form = GoodsReceiptForm(instance=goods_receipt)
    
    return render(request, 'dashboard/goods_receipts/form.html', {
        'form': form,
        'goods_receipt': goods_receipt,
        'page_title': 'Chỉnh sửa phiếu nhập hàng',
    })

@admin_required
def goods_receipt_detail(request, pk):
    """Chi tiết phiếu nhập hàng"""
    goods_receipt = get_object_or_404(
        GoodsReceipt.objects.select_related('purchase_order', 'purchase_order__supplier', 'warehouse', 'warehouse__store'),
        pk=pk
    )
    items = goods_receipt.items.select_related('product')
    
    return render(request, 'dashboard/goods_receipts/detail.html', {
        'goods_receipt': goods_receipt,
        'items': items,
        'page_title': f'Chi tiết phiếu nhập hàng: {goods_receipt.receipt_number}',
    })

@admin_required
def goods_receipt_confirm(request, pk):
    """Xác nhận phiếu nhập hàng"""
    goods_receipt = get_object_or_404(GoodsReceipt, pk=pk)
    
    if goods_receipt.status != 'draft':
        messages.error(request, 'Chỉ có thể xác nhận phiếu nhập ở trạng thái Nháp!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)

    if request.method == 'POST':
        for item in goods_receipt.items.all():
            product = item.product
            product.stock += item.quantity
            product.save()

        goods_receipt.status = 'confirmed'
        goods_receipt.save()

        purchase_order = goods_receipt.purchase_order
        if purchase_order:
            purchase_order.status = 'received'
            purchase_order.save()

        messages.success(request, f'Đã xác nhận phiếu nhập hàng {goods_receipt.receipt_number}!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)
    
    return render(request, 'dashboard/goods_receipts/confirm.html', {
        'goods_receipt': goods_receipt,
        'page_title': 'Xác nhận phiếu nhập hàng',
    })

@admin_required
def goods_receipt_cancel(request, pk):
    """Hủy phiếu nhập hàng"""
    goods_receipt = get_object_or_404(GoodsReceipt, pk=pk)
    
    if goods_receipt.status != 'draft':
        messages.error(request, 'Chỉ có thể hủy phiếu nhập ở trạng thái Nháp!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)

    if request.method == 'POST':
        goods_receipt.status = 'cancelled'
        goods_receipt.save()
        messages.success(request, f'Đã hủy phiếu nhập hàng {goods_receipt.receipt_number}!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)
    
    return render(request, 'dashboard/goods_receipts/cancel.html', {
        'goods_receipt': goods_receipt,
        'page_title': 'Hủy phiếu nhập hàng',
    })

@admin_required
def goods_receipt_delete(request, pk):
    """Xóa phiếu nhập hàng"""
    goods_receipt = get_object_or_404(GoodsReceipt, pk=pk)
    
    if goods_receipt.status != 'draft':
        messages.error(request, 'Chỉ có thể xóa phiếu nhập ở trạng thái Nháp!')
        return redirect('dashboard:goods_receipt_detail', pk=pk)
    
    if request.method == 'POST':
        goods_receipt.delete()
        messages.success(request, f'Đã xóa phiếu nhập hàng {goods_receipt.receipt_number}!')
        return redirect('dashboard:goods_receipt_list')
    
    return render(request, 'dashboard/goods_receipts/delete.html', {
        'goods_receipt': goods_receipt,
        'page_title': 'Xóa phiếu nhập hàng',
    })


# ==================== CUSTOMER GROUP MANAGEMENT ====================

@admin_required
def customer_group_list(request):
    """Danh sách nhóm khách hàng"""
    query = request.GET.get('q', '')
    customer_groups = CustomerGroup.objects.all()
    
    if query:
        customer_groups = customer_groups.filter(
            Q(name__icontains=query) |
            Q(description__icontains=query)
        )
    
    customer_groups = customer_groups.order_by('name')
    page_obj, paginator = paginate_queryset(customer_groups, request.GET.get('page'), 20)
    
    return render(request, 'dashboard/customer_groups/list.html', {
        'page_obj': page_obj,
        'query': query,
        'page_title': 'Nhóm khách hàng',
    })

@admin_required
def customer_group_create(request):
    """Tạo nhóm khách hàng mới"""
    if request.method == 'POST':
        form = CustomerGroupForm(request.POST)
        if form.is_valid():
            customer_group = form.save()
            messages.success(request, f'Đã tạo nhóm khách hàng {customer_group.name}!')
            return redirect('dashboard:customer_group_list')
    else:
        form = CustomerGroupForm()
    
    return render(request, 'dashboard/customer_groups/form.html', {
        'form': form,
        'page_title': 'Thêm nhóm khách hàng',
    })

@admin_required
def customer_group_edit(request, pk):
    """Chỉnh sửa nhóm khách hàng"""
    customer_group = get_object_or_404(CustomerGroup, pk=pk)
    
    if request.method == 'POST':
        form = CustomerGroupForm(request.POST, instance=customer_group)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật nhóm khách hàng {customer_group.name}!')
            return redirect('dashboard:customer_group_list')
    else:
        form = CustomerGroupForm(instance=customer_group)
    
    return render(request, 'dashboard/customer_groups/form.html', {
        'form': form,
        'customer_group': customer_group,
        'page_title': 'Chỉnh sửa nhóm khách hàng',
    })

@admin_required
def customer_group_detail(request, pk):
    """Chi tiết nhóm khách hàng"""
    customer_group = get_object_or_404(CustomerGroup, pk=pk)
    
    # Lấy khách hàng trong nhóm
    customers = Customer.objects.filter(
        customer_group=customer_group
    ).order_by('name')
    
    # Thống kê
    customer_count = customers.count()
    total_orders = Order.objects.filter(
        customer__customer_group=customer_group
    ).count()
    total_revenue = Order.objects.filter(
        customer__customer_group=customer_group,
        status='completed'
    ).aggregate(total=Sum('total_amount'))['total'] or 0
    
    return render(request, 'dashboard/customer_groups/detail.html', {
        'customer_group': customer_group,
        'customers': customers,
        'customer_count': customer_count,
        'total_orders': total_orders,
        'total_revenue': total_revenue,
        'page_title': f'Chi tiết nhóm khách hàng: {customer_group.name}',
    })

@admin_required
def customer_group_delete(request, pk):
    """Xóa nhóm khách hàng"""
    customer_group = get_object_or_404(CustomerGroup, pk=pk)
    
    if request.method == 'POST':
        # Kiểm tra xem có khách hàng nào không
        if Customer.objects.filter(customer_group=customer_group).exists():
            messages.error(request, 'Không thể xóa nhóm khách hàng đã có khách hàng!')
            return redirect('dashboard:customer_group_detail', pk=pk)
        
        customer_group.delete()
        messages.success(request, f'Đã xóa nhóm khách hàng {customer_group.name}!')
        return redirect('dashboard:customer_group_list')
    
    return render(request, 'dashboard/customer_groups/delete.html', {
        'customer_group': customer_group,
        'page_title': 'Xóa nhóm khách hàng',
    })


# ==================== STOCK MANAGEMENT ====================

@admin_required
def stock_movement_list(request):
    """Danh sách chuyển động kho"""
    query = request.GET.get('q', '')
    movement_type = request.GET.get('type', '')
    warehouse_filter = request.GET.get('warehouse', '')
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    
    movements = StockMovement.objects.select_related(
        'product', 'warehouse', 'warehouse__store', 'created_by', 'created_by__user'
    ).order_by('-created_at')
    
    if query:
        movements = movements.filter(
            Q(product__name__icontains=query) |
            Q(product__barcode__icontains=query) |
            Q(note__icontains=query)
        )
    
    if movement_type == 'in':
        movements = movements.filter(movement_type='import')
    elif movement_type == 'out':
        movements = movements.filter(movement_type='export')
    elif movement_type:
        movements = movements.filter(movement_type=movement_type)
    
    if warehouse_filter:
        movements = movements.filter(warehouse_id=warehouse_filter)
    
    if date_from:
        movements = movements.filter(created_at__date__gte=date_from)
    
    if date_to:
        movements = movements.filter(created_at__date__lte=date_to)
    
    page_obj, paginator = paginate_queryset(movements, request.GET.get('page'), 30)
    warehouses = Warehouse.objects.filter(store__is_active=True).select_related('store')
    
    return render(request, 'dashboard/stock_movements/list.html', {
        'page_obj': page_obj,
        'query': query,
        'warehouses': warehouses,
        'movement_type': movement_type,
        'warehouse_filter': warehouse_filter,
        'date_from': date_from,
        'date_to': date_to,
        'page_title': 'Chuyển động kho',
    })

@admin_required
def stock_balance_list(request):
    """Danh sách số dư tồn kho"""
    query = request.GET.get('q', '')
    warehouse_filter = request.GET.get('warehouse', '')
    low_stock_only = request.GET.get('low_stock', '')
    
    balances = StockBalance.objects.select_related('product', 'warehouse', 'warehouse__store').order_by('warehouse__store__name', 'product__name')
    
    if query:
        balances = balances.filter(
            Q(product__name__icontains=query) |
            Q(product__barcode__icontains=query)
        )
    
    if warehouse_filter:
        balances = balances.filter(warehouse_id=warehouse_filter)
    
    if low_stock_only:
        balances = balances.filter(quantity_on_hand__lte=F('product__min_stock'))
    
    page_obj, paginator = paginate_queryset(balances, request.GET.get('page'), 30)
    warehouses = Warehouse.objects.filter(store__is_active=True).select_related('store')
    
    return render(request, 'dashboard/stock_balances/list.html', {
        'page_obj': page_obj,
        'query': query,
        'warehouses': warehouses,
        'warehouse_filter': warehouse_filter,
        'low_stock_only': low_stock_only,
        'page_title': 'Số dư tồn kho',
    })

@admin_required
def stock_balance_detail(request, pk):
    """Chi tiết số dư tồn kho"""
    balance = get_object_or_404(
        StockBalance.objects.select_related('product', 'warehouse'), 
        pk=pk
    )
    
    # Lấy các chuyển động kho gần đây của sản phẩm này
    movements = StockMovement.objects.filter(
        product=balance.product,
        warehouse=balance.warehouse
    ).order_by('-created_at')[:20]
    
    return render(request, 'dashboard/stock_balances/detail.html', {
        'balance': balance,
        'movements': movements,
        'page_title': f'Chi tiết tồn kho: {balance.product.name}',
    })


# ==================== REPORT EXPORT ====================

@admin_required
def export_orders_excel(request):
    """Xuất báo cáo đơn hàng Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'Thư viện openpyxl chưa được cài đặt. Chạy: pip install openpyxl')
        return redirect('dashboard:orders_list')

    from io import BytesIO
    from django.utils import timezone

    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    status_filter = request.GET.get('status')

    orders = Order.objects.select_related('customer', 'store')

    if date_from:
        orders = orders.filter(created_at__date__gte=date_from)
    if date_to:
        orders = orders.filter(created_at__date__lte=date_to)
    if status_filter:
        orders = orders.filter(status=status_filter)

    orders = orders.order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Báo cáo đơn hàng'
    header_fill = PatternFill(start_color='4CAF50', end_color='4CAF50', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    headers = ['Mã đơn', 'Ngày đặt', 'Khách hàng', 'Cửa hàng', 'Tổng tiền', 'Trạng thái']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, order in enumerate(orders, 2):
        ws.cell(row=row_idx, column=1, value=order.order_number or order.order_id or f'#{order.pk}')
        ws.cell(row=row_idx, column=2, value=order.created_at.date())
        ws.cell(row=row_idx, column=2).number_format = 'DD/MM/YYYY'
        ws.cell(row=row_idx, column=3, value=order.customer.name if order.customer else 'Khách vãng lai')
        ws.cell(row=row_idx, column=4, value=order.store.name if order.store else 'Online')
        ws.cell(row=row_idx, column=5, value=float(order.total_amount))
        ws.cell(row=row_idx, column=5).number_format = '#,##0'
        ws.cell(row=row_idx, column=6, value=order.get_status_display())

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="bao_cao_don_hang_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response

@admin_required
def export_revenue_excel(request):
    """Xuất báo cáo doanh thu Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'Thư viện openpyxl chưa được cài đặt. Chạy: pip install openpyxl')
        return redirect('dashboard:index')

    from io import BytesIO
    from django.utils import timezone

    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    orders = Order.objects.filter(status='completed')

    if date_from:
        orders = orders.filter(created_at__date__gte=date_from)
    if date_to:
        orders = orders.filter(created_at__date__lte=date_to)

    revenue_by_date = list(
        orders.extra(
            select={'day': 'date(created_at)'}
        ).values('day').annotate(
            total_revenue=Sum('total_amount'),
            order_count=Count('id')
        ).order_by('day')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Báo cáo doanh thu'
    header_fill = PatternFill(start_color='2196F3', end_color='2196F3', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    bold_font = Font(bold=True)

    headers = ['Ngày', 'Số đơn hàng', 'Doanh thu', 'Trung bình/đơn']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, data in enumerate(revenue_by_date, 2):
        ws.cell(row=row_idx, column=1, value=data['day'])
        ws.cell(row=row_idx, column=1).number_format = 'DD/MM/YYYY'
        ws.cell(row=row_idx, column=2, value=data['order_count'])
        ws.cell(row=row_idx, column=3, value=float(data['total_revenue'] or 0))
        ws.cell(row=row_idx, column=3).number_format = '#,##0'
        avg_revenue = (data['total_revenue'] or 0) / data['order_count'] if data['order_count'] > 0 else 0
        ws.cell(row=row_idx, column=4, value=float(avg_revenue))
        ws.cell(row=row_idx, column=4).number_format = '#,##0'

    total_row = len(revenue_by_date) + 2
    ws.cell(row=total_row, column=1, value='TỔNG CỘNG')
    ws.cell(row=total_row, column=1).font = bold_font
    ws.cell(row=total_row, column=2, value=sum(item['order_count'] for item in revenue_by_date))
    total_rev = sum(item['total_revenue'] or 0 for item in revenue_by_date)
    ws.cell(row=total_row, column=3, value=float(total_rev))
    ws.cell(row=total_row, column=3).number_format = '#,##0'
    total_orders = sum(item['order_count'] for item in revenue_by_date)
    total_avg = total_rev / total_orders if total_orders > 0 else 0
    ws.cell(row=total_row, column=4, value=float(total_avg))
    ws.cell(row=total_row, column=4).number_format = '#,##0'

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="bao_cao_doanh_thu_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response

@admin_required
def export_stock_excel(request):
    """Xuất báo cáo tồn kho Excel"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'Thư viện openpyxl chưa được cài đặt. Chạy: pip install openpyxl')
        return redirect('dashboard:stock_balance_list')

    from io import BytesIO
    from django.utils import timezone

    warehouse_filter = request.GET.get('warehouse')
    low_stock_only = request.GET.get('low_stock')

    balances = StockBalance.objects.select_related('product', 'warehouse', 'warehouse__store')

    if warehouse_filter:
        balances = balances.filter(warehouse_id=warehouse_filter)
    if low_stock_only:
        balances = balances.filter(quantity_on_hand__lte=F('product__min_stock'))

    balances = balances.order_by('warehouse__store__name', 'product__name')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Báo cáo tồn kho'
    header_fill = PatternFill(start_color='FF9800', end_color='FF9800', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    headers = ['Kho', 'SKU', 'Sản phẩm', 'Tồn kho', 'Tồn tối thiểu', 'Trạng thái']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, balance in enumerate(balances, 2):
        ws.cell(row=row_idx, column=1, value=balance.warehouse.store.name)
        ws.cell(row=row_idx, column=2, value=balance.product.barcode or '')
        ws.cell(row=row_idx, column=3, value=balance.product.name)
        ws.cell(row=row_idx, column=4, value=balance.quantity_on_hand)
        ws.cell(row=row_idx, column=4).number_format = '#,##0'
        ws.cell(row=row_idx, column=5, value=balance.product.min_stock)
        ws.cell(row=row_idx, column=5).number_format = '#,##0'

        if balance.quantity_on_hand <= 0:
            status = 'Hết hàng'
        elif balance.quantity_on_hand <= balance.product.min_stock:
            status = 'Sắp hết'
        else:
            status = 'Tồn kho tốt'
        ws.cell(row=row_idx, column=6, value=status)

    widths = [20, 15, 30, 12, 12, 15]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="bao_cao_ton_kho_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


# ==================== DISCOUNT CODES ====================
@admin_required
def discount_code_list(request):
    """Danh sách mã giảm giá"""
    discount_codes = DiscountCode.objects.all().order_by('-created_at')
    
    # Search functionality
    query = request.GET.get('q', '')
    if query:
        discount_codes = discount_codes.filter(
            Q(code__icontains=query) | 
            Q(name__icontains=query) | 
            Q(description__icontains=query)
        )
    
    # Filter by status
    status_filter = request.GET.get('status', '')
    if status_filter:
        discount_codes = discount_codes.filter(status=status_filter)
    
    page_obj, paginator = paginate_queryset(discount_codes, request.GET.get('page'), 20)
    
    context = {
        'page_obj': page_obj,
        'paginator': paginator,
        'query': query,
        'status_filter': status_filter,
        'page_title': 'Quản Lý Mã Giảm Giá',
    }
    return render(request, 'dashboard/discount_codes/list.html', context)


@admin_required
def discount_code_create(request):
    """Tạo mã giảm giá mới"""
    if request.method == 'POST':
        form = DiscountCodeForm(request.POST)
        if form.is_valid():
            discount_code = form.save(commit=False)
            discount_code.created_by = request.user.employee if hasattr(request.user, 'employee') else None
            discount_code.save()
            messages.success(request, 'Mã giảm giá đã được tạo thành công!')
            return redirect('dashboard:discount_code_list')
    else:
        form = DiscountCodeForm()
    
    context = {
        'form': form,
        'page_title': 'Thêm Mã Giảm Giá',
        'is_create': True,
    }
    return render(request, 'dashboard/discount_codes/form.html', context)


@admin_required
def discount_code_edit(request, pk):
    """Chỉnh sửa mã giảm giá"""
    discount_code = get_object_or_404(DiscountCode, pk=pk)
    
    if request.method == 'POST':
        form = DiscountCodeForm(request.POST, instance=discount_code)
        if form.is_valid():
            form.save()
            messages.success(request, 'Mã giảm giá đã được cập nhật!')
            return redirect('dashboard:discount_code_list')
    else:
        form = DiscountCodeForm(instance=discount_code)
    
    context = {
        'form': form,
        'discount_code': discount_code,
        'page_title': f'Chỉnh sửa: {discount_code.code}',
        'is_edit': True,
    }
    return render(request, 'dashboard/discount_codes/form.html', context)


@admin_required
def discount_code_detail(request, pk):
    """Chi tiết mã giảm giá"""
    discount_code = get_object_or_404(DiscountCode, pk=pk)
    usages = DiscountCodeUsage.objects.select_related('user', 'order').filter(discount_code=discount_code).order_by('-used_at')[:20]
    
    # Calculate statistics
    total_usage = discount_code.usage_count
    total_discount = DiscountCodeUsage.objects.filter(discount_code=discount_code).aggregate(
        total=Sum('discount_amount')
    )['total'] or 0
    
    context = {
        'discount_code': discount_code,
        'usages': usages,
        'total_usage': total_usage,
        'total_discount': total_discount,
        'page_title': f'Chi tiết: {discount_code.code}',
    }
    return render(request, 'dashboard/discount_codes/detail.html', context)


@admin_required
@require_http_methods(["POST"])
def discount_code_delete(request, pk):
    """Xóa mã giảm giá"""
    discount_code = get_object_or_404(DiscountCode, pk=pk)
    code = discount_code.code
    
    # Check if code has been used
    if discount_code.usage_count > 0:
        messages.error(request, f'Không thể xóa mã "{code}" vì đã có {discount_code.usage_count} lượt sử dụng!')
        return redirect('dashboard:discount_code_detail', pk=pk)
    
    discount_code.delete()
    messages.success(request, f'Mã giảm giá "{code}" đã được xóa!')
    return redirect('dashboard:discount_code_list')


# ==================== DISCOUNT CODE API ====================
@api_view(['POST'])
@permission_classes([permissions.IsAuthenticated])
def validate_discount_code(request):
    """API validate và áp dụng mã giảm giá"""
    try:
        code = request.data.get('code', '').strip().upper()
        order_amount = float(request.data.get('order_amount', 0))
        
        if not code:
            return Response({'error': 'Vui lòng nhập mã giảm giá'}, status=400)
        
        if order_amount <= 0:
            return Response({'error': 'Giá trị đơn hàng không hợp lệ'}, status=400)
        
        # Find discount code
        try:
            discount_code = DiscountCode.objects.get(code=code)
        except DiscountCode.DoesNotExist:
            return Response({'error': 'Mã giảm giá không tồn tại'}, status=404)
        
        # Check if code is valid
        is_valid, message = discount_code.is_valid()
        if not is_valid:
            return Response({'error': message}, status=400)
        
        # Check if user can use this code
        can_use, user_message = discount_code.can_user_use(request.user)
        if not can_use:
            return Response({'error': user_message}, status=400)
        
        # Calculate discount
        discount_amount, error_message = discount_code.calculate_discount(order_amount)
        if error_message:
            return Response({'error': error_message}, status=400)
        
        # Return success response
        response_data = {
            'success': True,
            'discount_code': {
                'code': discount_code.code,
                'name': discount_code.name,
                'discount_type': discount_code.discount_type,
                'discount_value': float(discount_code.discount_value),
                'discount_amount': float(discount_amount),
                'final_amount': order_amount - discount_amount,
            }
        }
        
        return Response(response_data)
        
    except Exception as e:
        return Response({'error': f'Lỗi hệ thống: {str(e)}'}, status=500)
